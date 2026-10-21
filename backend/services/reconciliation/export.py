"""Excel file with the result of «наш расчёт ↔ файл сверки ЧДУ» — to keep or to
send to the manager with the list of differences."""
from __future__ import annotations

import io
from datetime import date
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="E2E8F0")
BAD = PatternFill("solid", fgColor="FDE2E1")
WARN = PatternFill("solid", fgColor="FEF3C7")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00;[Red]-#,##0.00'
STATUS = {"OK": "совпадает", "DIFF": "расхождение", "ONLY_OURS": "только у нас", "ONLY_CDU": "только у ЧДУ",
          "CLOSED": "закрыта на дату", "MISSING": "нет в файле"}


def _table(ws, start_row: int, headers: List[str], rows: List[List[Any]], formats: Dict[int, str],
           bad_rows: Dict[int, PatternFill]) -> int:
    for c, h in enumerate(headers, 1):
        cell = ws.cell(start_row, c, h)
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(start_row + i, c, v)
            cell.border = BORDER
            if c in formats and isinstance(v, (int, float)):
                cell.number_format = formats[c]
            if i - 1 in bad_rows:
                cell.fill = bad_rows[i - 1]
    return start_row + len(rows) + 2


def reconciliation_xlsx(d: Dict[str, Any], *, cdu_name: str, recon_date: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Итоги"
    ws["A1"] = f"Сверка портфеля КФГД с ЧДУ: {cdu_name}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Дата сверки: {recon_date:%d.%m.%Y}"
    f = d.get("file") or {}
    ws["A3"] = f"Файл ЧДУ: {f.get('file_name') or '—'}"
    ws["A4"] = "Итог: " + ("расхождений нет" if d.get("status") == "OK" else "есть расхождения")
    ws["A4"].font = Font(bold=True, color="166534" if d.get("status") == "OK" else "B91C1C")
    rows, bad = [], {}
    for i, s in enumerate(d.get("summary") or []):
        rows.append([s["label"], s.get("ours"), s.get("cdu"), s.get("delta"), STATUS.get(s["status"], s["status"])])
        if s["status"] != "OK":
            bad[i] = BAD
    r = _table(ws, 6, ["Раздел", "Наш расчёт, ₸", "Файл ЧДУ, ₸", "Разница, ₸", "Статус"], rows,
               {2: MONEY, 3: MONEY, 4: MONEY}, bad)
    for note in (d.get("notes") or []) + (d.get("file_warnings") or []):
        ws.cell(r, 1, note)
        r += 1
    for col, w in zip("ABCDE", (34, 22, 22, 18, 16)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Ценные бумаги")
    rows, bad = [], {}
    for i, s in enumerate(d.get("securities") or []):
        o, t = s.get("ours") or {}, s.get("cdu") or {}
        rows.append([s["isin"], s.get("ticker"), s.get("name"),
                     t.get("quantity"), o.get("quantity"), t.get("price"), o.get("price"),
                     t.get("accrued"), o.get("accrued"), t.get("value"), o.get("value"),
                     s.get("delta_value"), STATUS.get(s["status"], s["status"]), "; ".join(s.get("reasons") or [])])
        if s["status"] != "OK":
            bad[i] = BAD
    _table(ws2, 1, ["ISIN", "Тикер", "Наименование", "Кол-во ЧДУ", "Кол-во наше", "Цена ЧДУ", "Цена наша",
                    "НКД ЧДУ", "НКД наш", "Стоимость ЧДУ", "Стоимость наша", "Разница", "Статус", "Причина"],
           rows, {4: '#,##0', 5: '#,##0', 6: '0.0000', 7: '0.0000', 8: MONEY, 9: MONEY, 10: MONEY,
                  11: MONEY, 12: MONEY}, bad)
    for c, w in enumerate((15, 14, 34, 14, 14, 10, 10, 16, 16, 18, 18, 14, 14, 70), 1):
        ws2.column_dimensions[get_column_letter(c)].width = w
    ws2.freeze_panes = "B2"

    ws3 = wb.create_sheet("Обратное РЕПО")
    rows, bad = [], {}
    for i, s in enumerate(d.get("repos") or []):
        o, t = s.get("ours") or {}, s.get("cdu") or {}
        rows.append([s.get("isin"), t.get("open_date") or o.get("open_date"), t.get("close_date") or o.get("close_date"),
                     t.get("open_amount"), o.get("open_amount"), t.get("accrued"), o.get("accrued"),
                     t.get("value_effective"), o.get("value"), s.get("delta_value"),
                     STATUS.get(s["status"], s["status"]), "; ".join(s.get("reasons") or [])])
        if s["status"] not in ("OK", "CLOSED"):
            bad[i] = BAD
        elif s["status"] == "CLOSED":
            bad[i] = WARN
    _table(ws3, 1, ["Залог (ISIN)", "Открытие", "Закрытие", "Сумма открытия ЧДУ", "Сумма открытия наша",
                    "Вознаграждение ЧДУ", "Вознаграждение наше", "Стоимость ЧДУ", "Стоимость наша",
                    "Разница", "Статус", "Причина"],
           rows, {4: MONEY, 5: MONEY, 6: MONEY, 7: MONEY, 8: MONEY, 9: MONEY, 10: MONEY}, bad)
    for c, w in enumerate((15, 12, 12, 18, 18, 16, 16, 18, 18, 14, 16, 60), 1):
        ws3.column_dimensions[get_column_letter(c)].width = w
    ws3.freeze_panes = "B2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
