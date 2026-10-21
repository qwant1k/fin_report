"""Instrument sheets of the XLSX export: «ГЦБ», «Агентские», «МФО», «РЕПО»,
«Дебиторская задолженность».

The rows are built by the same portfolio engine as «Сводка» and its
drill-down (Risk Report slice rolled forward with the ledger), so the total
of a manager on a sheet equals the category value of the summary on the date.
On the date of a Risk Report the figures are those of the Risk Report.
"""
from __future__ import annotations

from datetime import date
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from models.db_models import CDU, PortfolioPosition, PortfolioSummary

HEAD_FILL = PatternFill("solid", fgColor="1F6B38")
CDU_FILL = PatternFill("solid", fgColor="E2EFDA")
TOTAL_FILL = PatternFill("solid", fgColor="C6E0B4")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
WHITE_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BOLD = Font(name="Calibri", size=11, bold=True)
MONEY = '# ##0.00;-# ##0.00;"-"'
QTY = '# ##0'
PCT = '0.00%'
PRICE = '0.0000'
DUR = '0.0000'
DATE = 'DD.MM.YYYY'

BOND_SHEETS = [
    ("ГЦБ", "GOV_BONDS"),
    ("Агентские", "AGENCY_BONDS"),
    ("МФО", "MFO_BONDS"),
]

BOND_COLUMNS = [
    ("ЧДУ", 34, None), ("Тикер", 14, None), ("ISIN", 15, None), ("Эмитент / наименование", 36, None),
    ("Дата сделки", 12, DATE), ("Дата расчётов", 12, DATE), ("Количество, шт", 14, QTY),
    ("Номинал (Face value)", 18, MONEY), ("Дата погашения", 12, DATE), ("Ставка купона", 10, PCT),
    ("Чистая цена, %", 12, PRICE), ("Источник цены", 26, None), ("НКД (Accrued interest)", 16, MONEY),
    ("Чистая стоимость (Market value)", 18, MONEY), ("Стоимость (Total value)", 18, MONEY),
    ("Вес", 9, PCT), ("YTM", 9, PCT), ("Mod. duration", 11, DUR),
    ("Взвеш. YTM", 10, PCT), ("Взвеш. duration", 11, DUR), ("Откуда позиция", 30, None),
]
REPO_COLUMNS = [
    ("ЧДУ", 34, None), ("Номер сделки", 14, None), ("Залог (тикер)", 14, None), ("ISIN залога", 15, None),
    ("Дата открытия", 12, DATE), ("Дата закрытия", 12, DATE), ("Дней прошло / всего", 12, None),
    ("Сумма открытия (Face value)", 18, MONEY), ("Ставка", 9, PCT), ("База", 7, None),
    ("Начислено (Accrued interest)", 16, MONEY), ("Стоимость (Market value)", 18, MONEY), ("Вес", 9, PCT),
    ("Сумма закрытия (Close value)", 18, MONEY), ("YTM", 9, PCT), ("Mod. duration", 11, DUR),
    ("Взвеш. YTM", 10, PCT), ("Взвеш. duration", 11, DUR), ("Откуда позиция", 30, None),
]
AR_COLUMNS = [
    ("№ п/п", 7, None), ("ЧДУ", 34, None), ("Контрагент", 30, None), ("Наименование / основание", 50, None),
    ("ISIN", 15, None), ("Валюта", 8, None), ("Дата постановки на учёт", 14, DATE), ("Срок погашения", 13, DATE),
    ("Первоначальная сумма, ₸", 18, MONEY), ("Погашено / скорректировано, ₸", 18, MONEY),
    ("Остаток в тенге", 18, MONEY), ("Откуда", 30, None),
]


def _d(v: Optional[str]) -> Optional[date]:
    return date.fromisoformat(v) if v else None


def _header(ws, title: str, report_date: date, columns) -> int:
    ws.cell(1, 1, title).font = Font(name="Calibri", size=13, bold=True)
    ws.cell(2, 1, "Дата отчёта")
    c = ws.cell(2, 2, report_date)
    c.number_format = DATE
    ws.cell(2, 3, "Источник: срез Risk Report + операции после него (как в «Сводке»)").font = Font(italic=True, color="666666")
    for i, (name, width, _) in enumerate(columns, start=1):
        cell = ws.cell(4, i, name)
        cell.fill = HEAD_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "C5"
    return 5


def _put(ws, row: int, values: list, columns, *, fill=None, font=None) -> None:
    for i, v in enumerate(values, start=1):
        cell = ws.cell(row, i, v)
        fmt = columns[i - 1][2]
        if fmt and v is not None and not isinstance(v, str):
            cell.number_format = fmt
        if fill is not None:
            cell.fill = fill
        if font is not None:
            cell.font = font


def _source_label(src: Optional[dict]) -> str:
    if not src:
        return ""
    parts = [src.get("label") or ""]
    if src.get("file_name"):
        parts.append(src["file_name"])
    return " · ".join(p for p in parts if p)


def _summary_values(db: Session, report_date: date) -> Dict[int, Dict[str, float]]:
    out: Dict[int, Dict[str, float]] = {}
    for p in db.execute(select(PortfolioPosition).where(
        PortfolioPosition.position_date == report_date,
        PortfolioPosition.instrument_code.is_(None),
    )).scalars():
        out.setdefault(p.cdu_id, {})[p.instrument_category] = float(p.market_value_current or 0.0)
    return out


def write_detail_sheets(wb: Workbook, db: Session, report_date: date) -> None:
    from services.portfolio.breakdown import _Sources, _bond_items, _money_market_items, _receivables
    from services.portfolio.engine import build_portfolio

    summaries = db.execute(select(PortfolioSummary).where(
        PortfolioSummary.summary_date == report_date,
    )).scalars().all()
    cdus = [db.get(CDU, s.cdu_id) for s in summaries]
    cdus = [c for c in cdus if c is not None]
    summary_values = _summary_values(db, report_date)
    src = _Sources(db)

    data: Dict[int, dict] = {}
    for cdu in cdus:
        state = build_portfolio(db, cdu.id, report_date)
        data[cdu.id] = {
            "bonds": {cat: _bond_items(db, src, state, cat) for _, cat in BOND_SHEETS},
            "repos": _money_market_items(src, state.repos, report_date, state.base_date, db),
            "ar": _receivables(db, src, cdu.id, report_date, state.base_date)["items"],
        }

    # ── ГЦБ / Агентские / МФО ──
    for title, cat in BOND_SHEETS:
        ws = wb.create_sheet(title)
        row = _header(ws, title, report_date, BOND_COLUMNS)
        grand = 0.0
        for cdu in cdus:
            items = data[cdu.id]["bonds"][cat]
            if not items:
                continue
            total = sum(i["value_kzt"] for i in items)
            grand += total
            for it in items:
                w = it["value_kzt"] / total if total else 0.0
                ytm, dur = it.get("ytm"), it.get("duration")
                name = " / ".join(x for x in (it.get("issuer"), it.get("name")) if x)
                _put(ws, row, [
                    cdu.name, it.get("code"), it.get("isin"), name,
                    _d(it.get("trade_date")), _d(it.get("settlement_date")), it.get("quantity"),
                    it.get("face_value"), _d(it.get("maturity_date")),
                    (it["coupon_rate_pct"] / 100.0) if it.get("coupon_rate_pct") is not None else None,
                    it.get("price"), it.get("price_source_label") or "", it.get("accrued"),
                    it.get("clean_value"), it["value_kzt"], w, ytm, dur,
                    (ytm or 0.0) * w if ytm is not None else None, (dur or 0.0) * w if dur is not None else None,
                    _source_label(it.get("source")),
                ], BOND_COLUMNS)
                row += 1
            _total_row(ws, row, cdu, total, summary_values.get(cdu.id, {}).get(cat), BOND_COLUMNS, value_col=15,
                       ytm=sum((i.get("ytm") or 0.0) * i["value_kzt"] for i in items if i.get("ytm") is not None) / total if total else None,
                       dur=sum((i.get("duration") or 0.0) * i["value_kzt"] for i in items if i.get("duration") is not None) / total if total else None,
                       wcols=(19, 20))
            row += 2
        _grand_row(ws, row, grand, BOND_COLUMNS, value_col=15)

    # ── РЕПО ──
    ws = wb.create_sheet("РЕПО")
    row = _header(ws, "Обратное РЕПО", report_date, REPO_COLUMNS)
    grand = 0.0
    for cdu in cdus:
        items = data[cdu.id]["repos"]
        if not items:
            continue
        total = sum(i["value_kzt"] for i in items)
        grand += total
        for it in items:
            w = it["value_kzt"] / total if total else 0.0
            ytm, dur = it.get("ytm"), it.get("duration")
            _put(ws, row, [
                cdu.name, it.get("deal_id") or it.get("code"), it.get("collateral_ticker") or it.get("collateral"),
                it.get("isin"), _d(it.get("open_date")), _d(it.get("close_date")),
                f"{it.get('days_passed') if it.get('days_passed') is not None else '—'} / "
                f"{it.get('days_total') if it.get('days_total') is not None else '—'}",
                it.get("open_amount"), it.get("rate"), it.get("base"), it.get("accrued"), it["value_kzt"], w,
                it.get("close_amount"), ytm, dur,
                (ytm or 0.0) * w if ytm is not None else None, (dur or 0.0) * w if dur is not None else None,
                _source_label(it.get("source")),
            ], REPO_COLUMNS)
            row += 1
        _total_row(ws, row, cdu, total, summary_values.get(cdu.id, {}).get("REVERSE_REPO"), REPO_COLUMNS,
                   value_col=12,
                   ytm=sum((i.get("ytm") or 0.0) * i["value_kzt"] for i in items if i.get("ytm") is not None) / total if total else None,
                   dur=sum((i.get("duration") or 0.0) * i["value_kzt"] for i in items if i.get("duration") is not None) / total if total else None,
                   wcols=(17, 18))
        row += 2
    _grand_row(ws, row, grand, REPO_COLUMNS, value_col=12)

    # ── Дебиторская задолженность ──
    ws = wb.create_sheet("Дебиторская задолженность")
    row = _header(ws, "Дебиторская задолженность", report_date, AR_COLUMNS)
    n = 0
    grand = 0.0
    for cdu in cdus:
        items = [i for i in data[cdu.id]["ar"] if abs(i["balance_kzt"]) > 0.005]
        if not items:
            continue
        total = sum(i["balance_kzt"] for i in items)
        grand += total
        for it in items:
            n += 1
            paid = sum(t["amount_kzt"] for t in it.get("transactions", []) if t.get("counted"))
            _put(ws, row, [
                n, cdu.name, it.get("counterparty"), it.get("description"), it.get("isin"), it.get("currency"),
                _d(it.get("record_date")), _d(it.get("due_date")), it.get("initial_kzt"), paid,
                it["balance_kzt"], _source_label(it.get("source")),
            ], AR_COLUMNS)
            row += 1
        _total_row(ws, row, cdu, total, summary_values.get(cdu.id, {}).get("RECEIVABLES"), AR_COLUMNS, value_col=11)
        row += 2
    if n == 0:
        ws.cell(row, 1, "Дебиторской задолженности на дату нет")
        row += 1
    _grand_row(ws, row, grand, AR_COLUMNS, value_col=11)


def _total_row(ws, row: int, cdu: CDU, total: float, summary_value: Optional[float], columns, *,
               value_col: int, ytm: Optional[float] = None, dur: Optional[float] = None,
               wcols: Optional[tuple] = None) -> None:
    values = [None] * len(columns)
    values[0] = f"Итого {cdu.short_name or cdu.name}"
    values[value_col - 1] = total
    if wcols and ytm is not None:
        values[wcols[0] - 1] = ytm
    if wcols and dur is not None:
        values[wcols[1] - 1] = dur
    _put(ws, row, values, columns, fill=CDU_FILL, font=BOLD)
    if summary_value is not None and abs(summary_value - total) > max(1.0, abs(summary_value) * 1e-6):
        cell = ws.cell(row, len(columns), f"В «Сводке» {summary_value:,.2f} — пересчитайте сводку за дату".replace(",", " "))
        cell.fill = WARN_FILL


def _grand_row(ws, row: int, total: float, columns, *, value_col: int) -> None:
    values = [None] * len(columns)
    values[0] = "ИТОГО по фонду"
    values[value_col - 1] = total
    _put(ws, row, values, columns, fill=TOTAL_FILL, font=BOLD)
