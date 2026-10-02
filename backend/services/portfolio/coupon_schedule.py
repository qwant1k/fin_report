"""Coupon payment schedules (Справочник ЦБ → «График купонных выплат»).

A schedule row is a payment: its date, the annual rate of the period that ends
on it and, optionally, the coupon per security.  When an issue has a schedule
it is the source of truth for coupon dates and rates (floating / indexed
coupons, irregular periods, dates shifted by holidays — e.g. KASE shows
«01.04.24» for MUM108_0013 instead of 31.03.24); otherwise the regular
schedule derived from the reference parameters is used.

Import accepts the table copied from kase.kz («№ | Дата начала купонной
выплаты | Ставка, % год | Дата фиксации реестра | Дата окончания выплаты») as
XLSX/CSV, optionally with an ISIN / ticker column for many issues at once.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

from sqlalchemy import select
from sqlalchemy.orm import Session

from models.db_models import CouponSchedule, InstrumentReference
from services.parser.number_utils import parse_kz_date, parse_kz_number
from services.portfolio import bond_math as bm


@dataclass
class ScheduleRow:
    payment_date: date
    rate: Optional[float]               # annual rate as a fraction
    amount_per_unit: Optional[float] = None
    period_start: Optional[date] = None
    record_date: Optional[date] = None
    payment_end_date: Optional[date] = None


def rows_for(db: Session, isin: str) -> List[ScheduleRow]:
    items = db.execute(select(CouponSchedule).where(
        CouponSchedule.isin == isin.strip().upper(),
    ).order_by(CouponSchedule.payment_date.asc())).scalars().all()
    return [_to_row(r) for r in items]


def load_all(db: Session) -> Dict[str, List[ScheduleRow]]:
    out: Dict[str, List[ScheduleRow]] = {}
    for r in db.execute(select(CouponSchedule).order_by(
        CouponSchedule.isin.asc(), CouponSchedule.payment_date.asc(),
    )).scalars():
        out.setdefault(r.isin.strip().upper(), []).append(_to_row(r))
    return out


def _to_row(r: CouponSchedule) -> ScheduleRow:
    return ScheduleRow(
        payment_date=r.payment_date, rate=bm.rate_fraction(r.rate_pct),
        amount_per_unit=r.amount_per_unit, period_start=r.period_start,
        record_date=r.record_date, payment_end_date=r.payment_end_date,
    )


# ───────────────────────── calculations ─────────────────────────

def period_around(schedule: List[ScheduleRow], as_of: date,
                  start: Optional[date] = None) -> Tuple[Optional[date], Optional[ScheduleRow]]:
    """(start of the current coupon period, the payment that closes it)."""
    prev: Optional[date] = None
    for row in schedule:
        if row.payment_date > as_of:
            begin = row.period_start or prev or start
            return begin, row
        prev = row.payment_date
    return prev, None


def accrued(face_value: float, schedule: List[ScheduleRow], as_of: date, basis: int,
            *, start: Optional[date] = None, nominal: Optional[float] = None) -> Optional[float]:
    """НКД on ``as_of`` from the schedule (None when the period is unknown)."""
    begin, row = period_around(schedule, as_of, start)
    if row is None or begin is None:
        return 0.0 if row is None and begin is not None else None
    if as_of <= begin:
        return 0.0
    if row.rate is not None:
        if basis == bm.BASIS_ACT_ACT:
            return face_value * row.rate * bm._year_fraction_act_act_split(begin, as_of)
        return face_value * row.rate * bm.year_fraction(begin, as_of, basis)
    if row.amount_per_unit is not None and nominal:
        full = bm.day_count(begin, row.payment_date, basis) or 1
        part = bm.day_count(begin, as_of, basis)
        return face_value / nominal * row.amount_per_unit * part / full
    return None


def coupon_on(schedule: List[ScheduleRow], d: date, face_value: float, basis: int,
              *, start: Optional[date] = None, nominal: Optional[float] = None
              ) -> Optional[Tuple[float, date, Optional[float]]]:
    """Coupon paid on ``d``: (amount, period start, rate) or None if ``d`` is not a payment date."""
    prev: Optional[date] = None
    for row in schedule:
        if row.payment_date == d:
            begin = row.period_start or prev or start
            if row.amount_per_unit is not None and nominal:
                return face_value / nominal * row.amount_per_unit, begin or d, row.rate
            if row.rate is None or begin is None:
                return None
            return bm.period_coupon_amount(face_value, row.rate, begin, d, basis), begin, row.rate
        if row.payment_date > d:
            return None
        prev = row.payment_date
    return None


def infer_frequency(schedule: List[ScheduleRow]) -> Optional[int]:
    dates = [r.payment_date for r in schedule]
    if len(dates) < 2:
        return None
    gaps = sorted((b - a).days for a, b in zip(dates, dates[1:]))
    median = gaps[len(gaps) // 2]
    for freq, days in ((12, 30), (4, 91), (2, 182), (1, 365)):
        if abs(median - days) <= days * 0.15:
            return freq
    return None


def schedule_mduration(as_of: date, schedule: List[ScheduleRow], maturity: Optional[date],
                       yld: float, basis: int, frequency: Optional[int] = None) -> Optional[float]:
    """Modified duration from the actual remaining cash flows of the schedule.

    Same conventions as Excel DURATION()/MDURATION() (time in coupon periods,
    the first one fractional, discounting at y/f per period), but every coupon
    uses the rate of its own period from the schedule — floating/indexed
    coupons and changed rates are taken into account; the principal is repaid
    with the last coupon (or at maturity if it is later).  For a schedule with
    a constant rate the result equals Excel MDURATION.
    """
    if yld is None:
        return None
    future = [r for r in schedule if r.payment_date > as_of]
    if not future:
        return None
    freq = frequency or infer_frequency(schedule) or 1
    y = yld / freq
    if 1 + y <= 0:
        return None
    past = [r.payment_date for r in schedule if r.payment_date <= as_of]
    ncd = future[0].payment_date
    pcd = past[-1] if past else None
    if pcd is None:
        from services.portfolio.bond_math import add_months
        pcd = add_months(ncd, -12 // freq if 12 % freq == 0 else -6)
    e = bm._coupdays(pcd, ncd, basis, freq)
    if e <= 0:
        return None
    frac = bm.day_count(as_of, ncd, basis) / e
    last_rate = next((r.rate for r in reversed(schedule) if r.rate is not None), None)
    flows: List[Tuple[float, float]] = []
    rate = None
    for k, row in enumerate(future, start=1):
        rate = row.rate if row.rate is not None else (rate if rate is not None else last_rate)
        cf = (rate or 0.0) * 100.0 / freq
        flows.append((k - 1 + frac, cf))
    t_last, cf_last = flows[-1]
    if maturity is not None and maturity > future[-1].payment_date:
        extra = (maturity - future[-1].payment_date).days / 365.0 * freq
        flows.append((t_last + extra, 100.0))
    else:
        flows[-1] = (t_last, cf_last + 100.0)
    num = sum(t * cf / (1 + y) ** t for t, cf in flows)
    den = sum(cf / (1 + y) ** t for t, cf in flows)
    if den == 0:
        return None
    return num / den / freq / (1 + y)


def current_rate(schedule: List[ScheduleRow], as_of: date) -> Optional[float]:
    _, row = period_around(schedule, as_of)
    return row.rate if row is not None else None


# ───────────────────────── maintenance ─────────────────────────

def generate_from_reference(ref: InstrumentReference) -> List[dict]:
    """Regular schedule from the issue parameters (start, maturity, frequency, rate)."""
    if not ref.maturity_date or not ref.frequency:
        return []
    dates = bm.coupon_schedule(ref.start_date, ref.maturity_date, int(ref.frequency))
    rate = ref.coupon_rate_pct
    out = []
    prev = ref.start_date
    for i, d in enumerate(dates, start=1):
        if ref.start_date and d <= ref.start_date:
            continue
        out.append({"seq": i, "period_start": prev, "payment_date": d, "rate_pct": rate})
        prev = d
    return out


def replace_schedule(db: Session, isin: str, rows: Iterable[dict], *, source: str,
                     actor: Optional[str], keep_manual: bool = False) -> int:
    """Replace the schedule of an issue.

    Each row may carry its own ``source`` (the UI marks edited/added rows as
    MANUAL).  With ``keep_manual`` (refresh from KASE) rows previously entered
    by hand win over the downloaded ones for the same payment date, and manual
    rows absent from the download are kept.
    """
    isin = isin.strip().upper()
    rows = list(rows)
    manual: Dict[date, CouponSchedule] = {}
    if keep_manual:
        for old in db.execute(select(CouponSchedule).where(
            CouponSchedule.isin == isin, CouponSchedule.source == "MANUAL",
        )).scalars():
            manual[old.payment_date] = old
    manual_rows = {d: {
        "seq": old.seq, "period_start": old.period_start, "payment_date": old.payment_date,
        "record_date": old.record_date, "payment_end_date": old.payment_end_date,
        "rate_pct": old.rate_pct, "margin_pct": old.margin_pct, "index_rate_pct": old.index_rate_pct,
        "amount_per_unit": old.amount_per_unit, "notes": old.notes, "source": "MANUAL",
        "updated_by": old.updated_by,
    } for d, old in manual.items()}
    db.query(CouponSchedule).filter(CouponSchedule.isin == isin).delete(synchronize_session="fetch")
    db.flush()
    merged: Dict[date, dict] = {}
    for r in rows:
        merged[r["payment_date"]] = {**r, "source": r.get("source") or source}
    merged.update(manual_rows)
    n = 0
    for i, r in enumerate(sorted(merged.values(), key=lambda x: x["payment_date"]), start=1):
        db.add(CouponSchedule(
            isin=isin, seq=i, period_start=r.get("period_start"),
            payment_date=r["payment_date"], record_date=r.get("record_date"),
            payment_end_date=r.get("payment_end_date"), rate_pct=r.get("rate_pct"),
            margin_pct=r.get("margin_pct"), index_rate_pct=r.get("index_rate_pct"),
            amount_per_unit=r.get("amount_per_unit"), source=str(r.get("source") or source)[:20],
            notes=r.get("notes"), updated_by=r.get("updated_by") or actor,
        ))
        n += 1
    rows = list(merged.values())
    ref = db.execute(select(InstrumentReference).where(InstrumentReference.isin == isin)).scalars().first()
    if ref is not None and n:
        today = date.today()
        upcoming = [r for r in rows if r["payment_date"] > today]
        nxt = min(upcoming, key=lambda r: r["payment_date"]) if upcoming else None
        if nxt is not None and nxt.get("rate_pct") is not None:
            ref.coupon_rate_pct = nxt["rate_pct"]
        ref.next_coupon_date = nxt["payment_date"] if nxt else ref.next_coupon_date
    db.flush()
    return n


_HEADERS = {
    "isin": ("isin", "нин", "код инструмента", "тикер", "ticker", "инструмент"),
    "seq": ("№", "no", "n", "номер"),
    "payment_date": ("дата начала купонной выплаты", "дата выплаты", "дата купона", "payment date",
                     "дата начала выплаты"),
    "rate_pct": ("ставка, % год", "ставка, %", "ставка", "rate", "купон, %"),
    "record_date": ("дата фиксации реестра", "фиксация реестра", "record date"),
    "payment_end_date": ("дата окончания выплаты", "окончание выплаты"),
    "amount_per_unit": ("сумма купона", "купон на одну бумагу", "сумма на 1 цб", "amount"),
    "period_start": ("начало периода", "дата начала периода"),
}


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip().lower().replace("ё", "е"))


def _match_header(cells: List) -> Optional[Dict[str, int]]:
    idx: Dict[str, int] = {}
    for i, c in enumerate(cells):
        t = _norm(c)
        if not t:
            continue
        for key, aliases in _HEADERS.items():
            if key in idx:
                continue
            if t in aliases or any(len(a) > 5 and a in t for a in aliases):
                idx[key] = i
                break
    return idx if "payment_date" in idx else None


def _date(v) -> Optional[date]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    text = str(v).strip()
    if text in ("-", "–", "—"):
        return None
    m = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{2})", text)
    if m:   # KASE shows two-digit years: 31.03.21
        return date(2000 + int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return parse_kz_date(text)


def parse_schedule_grid(grid: List[List]) -> Dict[Optional[str], List[dict]]:
    """Rows grouped by ISIN/ticker (None when the file has no instrument column)."""
    header_at, idx = None, None
    for r, row in enumerate(grid[:30]):
        idx = _match_header(list(row))
        if idx:
            header_at = r
            break
    if idx is None:
        raise ValueError("Не найдена колонка с датой выплаты (ожидается «Дата начала купонной выплаты» или «Дата выплаты»)")
    out: Dict[Optional[str], List[dict]] = {}
    for row in grid[header_at + 1:]:
        row = list(row)
        get = lambda k: row[idx[k]] if k in idx and idx[k] < len(row) else None  # noqa: E731
        pay = _date(get("payment_date"))
        if pay is None:
            continue
        rate = get("rate_pct")
        rate_v = parse_kz_number(rate) if not isinstance(rate, (int, float)) else float(rate)
        amount = get("amount_per_unit")
        amount_v = parse_kz_number(amount) if amount is not None and not isinstance(amount, (int, float)) else amount
        seq = get("seq")
        key = str(get("isin")).strip().upper() if get("isin") else None
        out.setdefault(key, []).append({
            "seq": int(float(seq)) if seq not in (None, "") and str(seq).replace(".", "", 1).isdigit() else None,
            "payment_date": pay,
            "rate_pct": rate_v,
            "record_date": _date(get("record_date")),
            "payment_end_date": _date(get("payment_end_date")),
            "amount_per_unit": float(amount_v) if amount_v not in (None, "") else None,
            "period_start": _date(get("period_start")),
        })
    return out


def read_grid(path: Path) -> List[List]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]
