"""KASE risk-parameter prices — "Расчетные цены_ФР".

The KASE app-gateway publishes a daily *risk-parameter* workbook (legacy
``.xls`` / OLE2) at::

    https://kase.kz/{lang}/app-gateway/documents/risk/{DD.MM.YYYY}

The workbook contains several sheets; the one we care about is
``Расчетные цены_ФР`` (settlement / "расчетные" prices for the stock market).
Layout (as published 2026)::

    row 0  : ["Дата оценки", …]
    row 1  : [<excel-serial valuation date>, …]
    row 3  : ["№ п/п", "Торговый код", "ISIN",
              "Алгоритмы и параметры расчета ЦБ",
              "Краткое наименование эмитента", "Расчетная цена",
              "Расчетная грязная", "Доходность до погашения. % годовых",
              "Дней до погашения", "Единица измерения цены"]
    row 4+ : data

These расчетные prices are used as a *fallback*: when an instrument has no
market price, its расчетная цена is applied instead.

Parsing is split into a thin xlrd loader (:func:`parse_risk_prices_xls`) and a
pure, dependency-free mapper (:func:`parse_risk_price_rows`) that is easy to
unit-test with synthetic rows.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional

import httpx
from loguru import logger

from config import settings
from services.kase.kase_client import _to_float

# Header label (normalised) → canonical field name.
_HEADER_EXACT: Dict[str, str] = {
    "торговый код": "instrument_code",
    "isin": "isin",
    "алгоритмы и параметры расчета цб": "algo",
    "краткое наименование эмитента": "issuer_name",
    "расчетная цена": "settlement_price",
    "расчетная грязная": "settlement_dirty_price",
    "единица измерения цены": "unit_ru",
}
# Long/variant labels matched by prefix.
_HEADER_PREFIX: tuple[tuple[str, str], ...] = (
    ("доходность", "dohod"),
    ("дней до погашения", "dtm"),
)
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0"
_XLSX_MAGIC = b"PK"


@dataclass
class KaseRiskQuote:
    instrument_code: str
    isin: Optional[str] = None
    algo: Optional[str] = None
    issuer_name: Optional[str] = None
    settlement_price: Optional[float] = None
    settlement_dirty_price: Optional[float] = None
    dohod: Optional[float] = None
    ytm: Optional[float] = None
    dtm: Optional[float] = None
    unit_ru: Optional[str] = None
    valuation_date: Optional[date] = None
    raw_data: Dict[str, object] = field(default_factory=dict)
    fetched_at: datetime = field(default_factory=datetime.utcnow)
    source: str = "risk_parameters"


def _norm(value: object) -> str:
    return str(value or "").strip().lower().replace("ё", "е")


def _excel_serial_to_date(serial: object, datemode: int = 0) -> Optional[date]:
    """Convert an Excel serial day number to a ``date``.

    ``datemode`` 0 → 1900 system (Windows, base 1899-12-30, which already
    accounts for the spurious 1900 leap day for any date after 1900-03-01);
    ``datemode`` 1 → 1904 system (legacy Mac).
    """
    if isinstance(serial, datetime):
        return serial.date()
    if isinstance(serial, date):
        return serial
    if isinstance(serial, str) and re.fullmatch(r"\s*\d{2}\.\d{2}\.\d{4}\s*", serial):
        return datetime.strptime(serial.strip(), "%d.%m.%Y").date()
    num = _to_float(serial)
    if num is None:
        return None
    base = date(1904, 1, 1) if datemode == 1 else date(1899, 12, 30)
    try:
        return base + timedelta(days=int(num))
    except (OverflowError, ValueError):
        return None


def _is_risk_price_sheet(name: str) -> bool:
    norm = _norm(name)
    # "Расчетные цены_ФР" — exclude the futures sheet "Расчетные цены_фьюч.".
    return "расчетные цены" in norm and "фьюч" not in norm


def _find_header_row(rows: List[list]) -> int:
    for idx, row in enumerate(rows):
        if any(_norm(cell) == "торговый код" for cell in row):
            return idx
    return -1


def _build_column_map(header: list) -> Dict[str, int]:
    col_map: Dict[str, int] = {}
    for col_idx, cell in enumerate(header):
        label = _norm(cell)
        if not label:
            continue
        if label in _HEADER_EXACT:
            col_map.setdefault(_HEADER_EXACT[label], col_idx)
            continue
        for prefix, fieldname in _HEADER_PREFIX:
            if label.startswith(prefix):
                col_map.setdefault(fieldname, col_idx)
                break
    return col_map


def _find_valuation_date(rows: List[list], header_idx: int, datemode: int) -> Optional[date]:
    for r in range(min(header_idx, len(rows))):
        for c, cell in enumerate(rows[r]):
            if _norm(cell) == "дата оценки":
                # value sits in the same column, on a row below the label
                for rr in range(r + 1, header_idx):
                    if c < len(rows[rr]):
                        found = _excel_serial_to_date(rows[rr][c], datemode)
                        if found:
                            return found
    return None


def _cell(row: list, col: Optional[int]):
    if col is None or col >= len(row):
        return None
    return row[col]


def _clean_str(value: object) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = text.strip()
    return text or None


def parse_risk_price_rows(
    rows: List[list],
    *,
    datemode: int = 0,
    fallback_date: Optional[date] = None,
) -> List[KaseRiskQuote]:
    """Map raw ``Расчетные цены_ФР`` rows to :class:`KaseRiskQuote` objects.

    ``rows`` is a list of cell-value lists (str / float / "") exactly as
    returned by xlrd's ``row_values``. Pure function — no I/O, no xlrd.
    """
    if not rows:
        return []
    header_idx = _find_header_row(rows)
    if header_idx < 0:
        logger.warning("Расчетные цены_ФР: header row ('Торговый код') not found")
        return []
    header = rows[header_idx]
    col_map = _build_column_map(header)
    code_col = col_map.get("instrument_code")
    price_col = col_map.get("settlement_price")
    if code_col is None or price_col is None:
        logger.warning(f"Расчетные цены_ФР: required columns missing, map={col_map}")
        return []

    headers_by_col = {c: _clean_str(header[c]) or f"col{c}" for c in range(len(header))}
    valuation_date = _find_valuation_date(rows, header_idx, datemode) or fallback_date

    out: List[KaseRiskQuote] = []
    for row in rows[header_idx + 1:]:
        code = _clean_str(_cell(row, code_col))
        if not code or _norm(code) == "торговый код":
            continue
        dohod = _to_float(_cell(row, col_map.get("dohod")))
        raw = {
            headers_by_col.get(c, f"col{c}"): row[c]
            for c in range(len(row))
            if c < len(row) and row[c] not in (None, "")
        }
        out.append(KaseRiskQuote(
            instrument_code=code,
            isin=_clean_str(_cell(row, col_map.get("isin"))),
            algo=_clean_str(_cell(row, col_map.get("algo"))),
            issuer_name=_clean_str(_cell(row, col_map.get("issuer_name"))),
            settlement_price=_to_float(_cell(row, price_col)),
            settlement_dirty_price=_to_float(_cell(row, col_map.get("settlement_dirty_price"))),
            dohod=dohod,
            ytm=(dohod / 100.0) if dohod is not None else None,
            dtm=_to_float(_cell(row, col_map.get("dtm"))),
            unit_ru=_clean_str(_cell(row, col_map.get("unit_ru"))),
            valuation_date=valuation_date,
            raw_data=raw,
        ))
    return out


def parse_risk_prices_xls(
    content: bytes,
    *,
    fallback_date: Optional[date] = None,
) -> List[KaseRiskQuote]:
    """Open a KASE risk-parameter workbook (``.xls`` or ``.xlsx``) and parse
    ``Расчетные цены_ФР``.  xlrd 2.x reads only legacy ``.xls``; ``.xlsx``
    files (KASE publishes both) are read with openpyxl."""
    if content[:2] == _XLSX_MAGIC:
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            name = next((n for n in wb.sheetnames if _is_risk_price_sheet(n)), None)
            if name is None:
                logger.warning(f"'Расчетные цены_ФР' sheet not found; available: {wb.sheetnames}")
                return []
            rows = [["" if v is None else v for v in r] for r in wb[name].iter_rows(values_only=True)]
        finally:
            wb.close()
        return parse_risk_price_rows(rows, datemode=0, fallback_date=fallback_date)

    import xlrd  # lazy: keeps the package importable without xlrd installed

    book = xlrd.open_workbook(file_contents=content)
    sheet_name = next((n for n in book.sheet_names() if _is_risk_price_sheet(n)), None)
    if sheet_name is None:
        logger.warning(
            f"'Расчетные цены_ФР' sheet not found; available: {book.sheet_names()}"
        )
        return []
    sheet = book.sheet_by_name(sheet_name)
    rows = [sheet.row_values(r) for r in range(sheet.nrows)]
    return parse_risk_price_rows(rows, datemode=book.datemode, fallback_date=fallback_date)


class KaseRiskClient:
    """Async downloader for the KASE risk-parameter "Расчетные цены_ФР" table."""

    def __init__(self) -> None:
        self.base_url = settings.kase_risk_params_url.rstrip("/")
        self._headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                          "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
            "Accept-Language": "ru,en;q=0.8",
            "Referer": "https://kase.kz/ru/markets/markets-valuation/risk-parameters",
        }

    def url_for(self, trade_date: date) -> str:
        return f"{self.base_url}/{trade_date.strftime('%d.%m.%Y')}"

    async def fetch_risk_prices(self, trade_date: date) -> List[KaseRiskQuote]:
        """Download + parse the расчетные цены for ``trade_date``.

        Returns an empty list when the file is missing for that date (KASE
        returns 404 on non-business days).
        """
        url = self.url_for(trade_date)
        async with httpx.AsyncClient(timeout=60.0, headers=self._headers, follow_redirects=True) as cli:
            try:
                resp = await cli.get(url)
            except Exception as exc:  # noqa: BLE001
                logger.warning(f"KASE risk-params fetch failed for {url}: {exc!r}")
                return []
            if resp.status_code == 404:
                logger.info(f"KASE risk-params: no file for {trade_date.isoformat()} ({url})")
                return []
            resp.raise_for_status()
            content = resp.content

        if not content:
            return []
        if not content.startswith(_OLE2_MAGIC) and not content.startswith(_XLSX_MAGIC):
            logger.warning(
                f"KASE risk-params: unexpected payload for {trade_date.isoformat()} "
                f"(first bytes {content[:8]!r})"
            )
            return []
        quotes = parse_risk_prices_xls(content, fallback_date=trade_date)
        logger.info(f"KASE risk-params {trade_date.isoformat()}: parsed {len(quotes)} rows")
        return quotes
