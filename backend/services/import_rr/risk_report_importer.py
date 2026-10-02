"""Главный оркестратор импорта Risk Report XLSM в БД.

Идемпотентность: для каждой пары (cdu_id, snapshot_date, ...) делается upsert,
повторный запуск за ту же дату не создаёт дубликатов (соответствует требованию
бизнес-процесса 5.4 "Про повторный запуск").
"""
from __future__ import annotations

import json
from io import BytesIO
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from zipfile import BadZipFile
from typing import Iterable, Optional

import msoffcrypto
from msoffcrypto.exceptions import DecryptionError, InvalidKeyError
from loguru import logger
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from config import settings
from models.db_models import (
    AccountReceivable,
    BondLot,
    CashSnapshot,
    CDU,
    DepositLot,
    FXRate,
    ImportJob,
    InstrumentReference,
    MBMIndex,
    MVSnapshot,
    PortfolioPosition,
    PortfolioSummary,
    RepoLot,
    SourceDocument,
)
from services.calculator.portfolio_calculator import calculate_for_date

from .helpers import (
    extract_date_from_filename,
    file_sha256,
    materialize_sheet,
    safe_sheet,
)


def _mat(wb, *names):
    """Найти лист в книге и сразу материализовать его в in-memory снимок.

    Решает проблему стейт-блида openpyxl read_only при многократном
    ``ws.cell()`` поверх нескольких листов одной книги.
    """
    ws = safe_sheet(wb, *names)
    return materialize_sheet(ws) if ws is not None else None


def _sheet_key(name: str) -> str:
    return "".join(ch for ch in str(name).lower() if ch.isalnum())


def _mat_exact(wb, *names):
    targets = {_sheet_key(name) for name in names}
    for sheet_name in wb.sheetnames:
        if _sheet_key(sheet_name) in targets:
            return materialize_sheet(wb[sheet_name])
    return None


def _load_risk_workbook(fp: Path, password: Optional[str] = None):
    """Open normal OOXML workbooks and password-encrypted Office files.

    Some SIUA files have an ``.xlsm`` extension but are actually OLE containers
    with ``EncryptedPackage``. ``openpyxl`` cannot open those until they are
    decrypted, so we use ``msoffcrypto`` first when needed.
    """
    try:
        return load_workbook(fp, read_only=True, data_only=True, keep_vba=False)
    except (BadZipFile, OSError) as first_exc:
        try:
            with fp.open("rb") as src:
                office = msoffcrypto.OfficeFile(src)
                if not office.is_encrypted():
                    raise first_exc

                pwd = password or settings.risk_report_password or None
                if not pwd:
                    raise ValueError(
                        "Файл Risk Report зашифрован. Укажите пароль при загрузке "
                        "или задайте RISK_REPORT_PASSWORD в .env."
                    )

                decrypted = BytesIO()
                office.load_key(password=pwd, verify_password=True)
                office.decrypt(decrypted)
                decrypted.seek(0)
        except (DecryptionError, InvalidKeyError) as exc:
            raise ValueError(
                "Пароль неверный. Введите правильный пароль для Risk Report и повторите загрузку."
            ) from exc
        except Exception:
            raise

        wb = load_workbook(decrypted, read_only=True, data_only=True, keep_vba=False)
        wb._decrypted_stream = decrypted  # keep BytesIO alive for read_only mode
        return wb
from .sheet_parsers import (
    get_report_date,
    parse_ar_sheet,
    parse_bond_lots_sheet,
    parse_cash_sheet,
    parse_dep_sheet,
    parse_fx_sheet,
    parse_mbm_sheet,
    parse_mv_sheet,
    parse_report_sheet,
    parse_reference_sheet,
    parse_repo_sheet,
)


@dataclass
class ImportResult:
    """Сводка по результату одного файла."""
    file_path: Path
    file_date: Optional[date] = None
    source_doc_id: Optional[int] = None
    cash_rows: int = 0
    mv_rows: int = 0
    reference_rows: int = 0
    fx_rows: int = 0
    mbm_rows: int = 0
    bond_lots: int = 0
    repo_lots: int = 0
    deposit_lots: int = 0
    ar_rows: int = 0
    report_summaries: int = 0
    report_positions: int = 0
    report_dates: list[date] = field(default_factory=list)
    portfolio_cdus: int = 0
    portfolio_breaches: int = 0
    skipped: bool = False
    error: Optional[str] = None
    warnings: list[str] = field(default_factory=list)

    def total_rows(self) -> int:
        return (self.cash_rows + self.mv_rows + self.reference_rows + self.fx_rows
                + self.mbm_rows + self.bond_lots + self.repo_lots
                + self.deposit_lots + self.ar_rows
                + self.report_summaries + self.report_positions)

    def to_log_line(self) -> str:
        if self.error:
            return f"[ERR ] {self.file_path.name}: {self.error}"
        if self.skipped:
            return f"[SKIP] {self.file_path.name}: {self.warnings[0] if self.warnings else 'duplicate'}"
        return (f"[OK  ] {self.file_path.name} ({self.file_date}): "
                f"cash={self.cash_rows} mv={self.mv_rows} ref={self.reference_rows} "
                f"fx={self.fx_rows} mbm={self.mbm_rows} bonds={self.bond_lots} "
                f"repo={self.repo_lots} dep={self.deposit_lots} ar={self.ar_rows} "
                f"report={self.report_summaries}/{self.report_positions} "
                f"portfolio_cdus={self.portfolio_cdus}")


def _refresh_portfolio_after_import(db: Session, res: ImportResult, report_date: Optional[date]) -> None:
    if report_date is None:
        return
    db.flush()
    counters = calculate_for_date(db, report_date, recalculate=True)
    res.portfolio_cdus = max(res.portfolio_cdus, int(counters.get("cdus_processed") or 0))
    res.portfolio_breaches += int(counters.get("breaches") or 0)


def _meta_report_t1_checked(doc: SourceDocument) -> bool:
    try:
        meta = json.loads(doc.parse_meta_json or "{}")
    except (TypeError, ValueError):
        return False
    dates = meta.get("report_dates") or []
    return bool(meta.get("report_t1_checked")) or len(set(dates)) >= 2


def _meta_report_dates(doc: SourceDocument) -> list[date]:
    try:
        meta = json.loads(doc.parse_meta_json or "{}")
    except (TypeError, ValueError):
        meta = {}
    out: list[date] = []
    for raw in meta.get("report_dates") or []:
        try:
            out.append(date.fromisoformat(str(raw)))
        except ValueError:
            continue
    if not out and doc.doc_date:
        out.append(doc.doc_date)
    return out


def _load_cdu_map(db: Session) -> dict[str, int]:
    """Current and canonical names → id (managers may be renamed in settings)."""
    from services.cdu_lookup import cdu_name_map
    cdu_map = cdu_name_map(db)
    if cdu_map:
        return cdu_map

    from seed import seed_initial_data

    seed_initial_data(db)
    return cdu_name_map(db)


# ═══════════════════════════════════════════════════════════════════════════
# Single-file import
# ═══════════════════════════════════════════════════════════════════════════
def import_risk_report(
    db: Session,
    file_path: Path | str,
    *,
    uploaded_by: Optional[str] = None,
    skip_if_imported: bool = False,
    password: Optional[str] = None,
) -> ImportResult:
    """Импорт одного XLSM-файла Risk Report в БД.

    Возвращает `ImportResult`. Все исключения ловит и записывает в `error`.
    """
    fp = Path(file_path)
    res = ImportResult(file_path=fp)

    try:
        # 1. Дата отчёта
        file_date = extract_date_from_filename(fp)

        # 2. Дедупликация по sha256 (если такой файл уже импортирован — пропустить)
        sha = file_sha256(fp) if fp.exists() else None
        if skip_if_imported and sha:
            existing = db.execute(select(SourceDocument).where(
                SourceDocument.sha256 == sha,
                SourceDocument.parse_status == "OK",
            ).order_by(SourceDocument.id.desc())).scalars().first()
            if existing:
                if _meta_report_t1_checked(existing):
                    res.skipped = True
                    res.warnings.append(f"уже импортирован, source_doc_id={existing.id}")
                    res.source_doc_id = existing.id
                    res.file_date = existing.doc_date
                    res.report_dates = _meta_report_dates(existing)
                    for refresh_date in res.report_dates:
                        _refresh_portfolio_after_import(db, res, refresh_date)
                    return res
                res.warnings.append(
                    f"повторный импорт: старый source_doc_id={existing.id} без Report (T-1)"
                )

        # 3. Создать SourceDocument запись
        sd = SourceDocument(
            doc_type="RISK_REPORT_XLSM",
            cdu_id=None,
            doc_date=file_date,
            file_name=fp.name,
            file_path=str(fp.resolve()),
            sha256=sha,
            file_size=fp.stat().st_size if fp.exists() else None,
            uploaded_by=uploaded_by,
            parse_status="PENDING",
        )
        db.add(sd)
        db.flush()
        res.source_doc_id = sd.id

        # 4. Открыть книгу (read_only + data_only — формулы как значения)
        wb = _load_risk_workbook(fp, password=password)

        # 5. Дата отчёта: Report!B2 — авторитетный источник (это «дата снимка»),
        # имя файла часто содержит дату генерации, а не отчёта (например,
        # risk_report_09112025_.xlsm может содержать снимок на 20.10.2025).
        report_ws = _mat(wb, "Report")
        report_b2 = get_report_date(report_ws) if report_ws else None
        if report_b2:
            file_date = report_b2
        elif file_date is None:
            file_date = date.today()
        res.file_date = file_date
        sd.doc_date = file_date

        # 6. Кеш: имя CDU → id
        cdu_map = _load_cdu_map(db)

        # 7. Парсеры по листам
        report_dates: set[date] = set()

        def _import_report_dashboard(ws, fallback_date: Optional[date]) -> tuple[int, int]:
            rows = parse_report_sheet(ws, fallback_date=fallback_date)
            for row in rows:
                snapshot_date = row.get("snapshot_date")
                if snapshot_date:
                    report_dates.add(snapshot_date)
            return _upsert_report_dashboard(db, rows, cdu_map)

        if report_ws is not None:
            summaries, positions = _import_report_dashboard(report_ws, file_date)
            res.report_summaries += summaries
            res.report_positions += positions

        report_t1_ws = _mat_exact(
            wb,
            "Report (T-1)",
            "Report (t-1)",
            "Report T-1",
            "Report t-1",
            "Report T1",
            "Report t1",
            "Report_T-1",
            "Report_t-1",
        )
        report_t1_checked = report_t1_ws is not None
        if report_t1_ws is not None:
            report_t1_date = get_report_date(report_t1_ws)
            fallback_t1_date = report_t1_date or (file_date - timedelta(days=1) if file_date else None)
            summaries, positions = _import_report_dashboard(
                report_t1_ws,
                fallback_t1_date,
            )
            res.report_summaries += summaries
            res.report_positions += positions

        res.report_dates = sorted(report_dates)

        # 7.1 Cash
        cash_ws = _mat(wb, "Cash")
        if cash_ws is not None:
            rows = parse_cash_sheet(cash_ws, fallback_date=file_date)
            res.cash_rows = _upsert_cash(db, rows, cdu_map, sd.id)

        # 7.2 MV
        mv_ws = _mat(wb, "MV")
        if mv_ws is not None:
            rows = parse_mv_sheet(mv_ws, fallback_date=file_date)
            res.mv_rows = _upsert_mv(db, rows, cdu_map, sd.id)

        # 7.3 Справочник
        ref_ws = _mat(wb, "Справочник", "Spravochnik", "Reference")
        if ref_ws is not None:
            rows = parse_reference_sheet(ref_ws)
            res.reference_rows = _upsert_reference(db, rows)

        # 7.4 FX
        fx_ws = _mat(wb, "Нацбанк Казахстана, Доллар США_",
                     "Нацбанк", "Доллар США", "USD KZT")
        if fx_ws is not None:
            rows = parse_fx_sheet(fx_ws)
            res.fx_rows = _upsert_fx(db, rows)

        # 7.5 MBM (история)
        mbm_ws = _mat(wb, "MBM index - с 1 апреля 2024 г",
                      "MBM index", "MBM Index")
        if mbm_ws is not None:
            rows = parse_mbm_sheet(mbm_ws)
            res.mbm_rows = _upsert_mbm(db, rows)

        _delete_rr_detail_rows_for_dates(db, [file_date])

        # 7.6 Bond lots
        for sheet_name, category in [
            ("ГЦБ",        "GOV_BONDS"),
            ("Агентские",  "AGENCY_BONDS"),
            ("Агентсткие", "AGENCY_BONDS"),  # частая опечатка в реальных файлах
            ("МФО",        "MFO_BONDS"),
            ("Ин. ЦБ",     "FOREIGN_BONDS"),
            ("Ин.ЦБ",      "FOREIGN_BONDS"),
            ("In CB",      "FOREIGN_BONDS"),
        ]:
            ws = _mat(wb, sheet_name)
            if ws is None:
                continue
            rows = parse_bond_lots_sheet(ws, category=category, fallback_date=file_date)
            if rows:
                res.bond_lots += _upsert_bond_lots(db, rows, cdu_map, sd.id, file_date)

        # 7.7 REPO lots
        repo_ws = _mat(wb, "Repo")
        if repo_ws is not None:
            rows = parse_repo_sheet(repo_ws, fallback_date=file_date)
            res.repo_lots = _upsert_repo_lots(db, rows, cdu_map, sd.id, file_date)

        # 7.8 Deposit lots
        dep_ws = _mat(wb, "Dep", "Депозит")
        if dep_ws is not None:
            rows = parse_dep_sheet(dep_ws, fallback_date=file_date)
            res.deposit_lots = _upsert_deposit_lots(db, rows, cdu_map, sd.id, file_date)

        # 7.9 AR
        ar_ws = _mat(wb, "Accounts receivable", "Дебиторская задолженность",
                     "Дебиторка")
        if ar_ws is not None:
            rows = parse_ar_sheet(ar_ws, fallback_date=file_date)
            res.ar_rows = _upsert_ar(db, rows, cdu_map, sd.id, file_date)

        # 8. Финализация
        wb.close()
        sd.parsed_at = datetime.utcnow()
        sd.parse_status = "OK"
        sd.rows_imported = res.total_rows()
        parse_meta = {
            "cash": res.cash_rows, "mv": res.mv_rows, "ref": res.reference_rows,
            "fx": res.fx_rows, "mbm": res.mbm_rows, "bonds": res.bond_lots,
            "repo": res.repo_lots, "dep": res.deposit_lots, "ar": res.ar_rows,
            "report_summaries": res.report_summaries,
            "report_positions": res.report_positions,
            "report_dates": [d.isoformat() for d in res.report_dates],
            "report_t1_checked": report_t1_checked,
        }
        sd.parse_meta_json = json.dumps(parse_meta, ensure_ascii=False)
        refresh_dates = res.report_dates or ([file_date] if file_date else [])
        for refresh_date in refresh_dates:
            _refresh_portfolio_after_import(db, res, refresh_date)
        parse_meta["portfolio_cdus"] = res.portfolio_cdus
        parse_meta["portfolio_breaches"] = res.portfolio_breaches
        sd.parse_meta_json = json.dumps(parse_meta, ensure_ascii=False)
        db.commit()

    except Exception as exc:
        logger.exception(f"import_risk_report failed for {fp}")
        db.rollback()
        res.error = str(exc) or repr(exc)
        # Best-effort: пометить SourceDocument как ERROR (уже создан)
        try:
            if res.source_doc_id:
                sd2 = db.get(SourceDocument, res.source_doc_id)
                if sd2:
                    sd2.parse_status = "ERROR"
                    sd2.parse_errors = (str(exc) or repr(exc))[:1000]
                    sd2.parsed_at = datetime.utcnow()
                    db.commit()
        except Exception:
            pass

    return res


# ═══════════════════════════════════════════════════════════════════════════
# Bulk folder import
# ═══════════════════════════════════════════════════════════════════════════
def import_folder(
    db: Session,
    folder: Path | str,
    *,
    uploaded_by: Optional[str] = None,
    job: Optional[ImportJob] = None,
    pattern: str = "**/*.xlsm",
    password: Optional[str] = None,
    skip_if_imported: bool = False,
) -> ImportJob:
    """Рекурсивно импортирует все XLSM-файлы из папки.

    Если `job` не передан, создаётся новый. Прогресс пишется в `job.files_done` и т.д.
    """
    folder_p = Path(folder)
    files: list[Path] = sorted(folder_p.glob(pattern))

    if job is None:
        job = ImportJob(
            job_type="HISTORIC_RR",
            status="RUNNING",
            triggered_by=uploaded_by,
            files_total=len(files),
            params_json=json.dumps({
                "folder": str(folder_p),
                "pattern": pattern,
                "skip_if_imported": skip_if_imported,
            }),
        )
        db.add(job)
        db.commit()
        db.refresh(job)
    else:
        job.files_total = len(files)
        job.status = "RUNNING"
        db.commit()

    log_lines: list[str] = []
    total_rows = 0
    failed = 0

    for i, fp in enumerate(files, start=1):
        # Свежая сессия не нужна — мы внутри Session, но коммиты атомарны
        result = import_risk_report(
            db,
            fp,
            uploaded_by=uploaded_by,
            password=password,
            skip_if_imported=skip_if_imported,
        )
        log_lines.append(result.to_log_line())
        if result.error:
            failed += 1
        else:
            total_rows += result.total_rows()

        # Обновлять прогресс каждые 5 файлов или в конце
        if i % 5 == 0 or i == len(files):
            job.files_done = i
            job.files_failed = failed
            job.rows_imported = total_rows
            job.log = "\n".join(log_lines[-200:])  # последние 200 строк
            db.commit()

    job.finished_at = datetime.utcnow()
    job.status = "DONE" if failed == 0 else ("PARTIAL" if failed < len(files) else "FAILED")
    job.files_done = len(files)
    job.files_failed = failed
    job.rows_imported = total_rows
    job.log = "\n".join(log_lines)
    db.commit()
    db.refresh(job)
    return job


# ═══════════════════════════════════════════════════════════════════════════
# Internal upsert helpers
# ═══════════════════════════════════════════════════════════════════════════
def _upsert_report_dashboard(
    db: Session, rows: list[dict], cdu_map: dict[str, int],
) -> tuple[int, int]:
    summary_rows = [r for r in rows if r.get("kind") == "summary"]
    position_rows = [r for r in rows if r.get("kind") == "position"]

    summaries_by_key: dict[tuple[int, date], PortfolioSummary] = {}
    fund_totals_by_date: dict[date, float] = {}
    summary_count = 0
    position_count = 0
    affected_dates = {
        r["snapshot_date"]
        for r in rows
        if r.get("snapshot_date")
    }
    if affected_dates:
        db.query(PortfolioPosition).filter(
            PortfolioPosition.position_date.in_(affected_dates),
            PortfolioPosition.notes == "rr_report_import",
        ).delete(synchronize_session=False)
        db.query(PortfolioSummary).filter(
            PortfolioSummary.summary_date.in_(affected_dates),
        ).delete(synchronize_session=False)

    def _summary_quality(values: dict) -> tuple[int, int]:
        return (
            1 if values.get("ytm_weighted") else 0,
            1 if values.get("duration_weighted") else 0,
        )

    for r in summary_rows:
        cdu_id = _resolve_cdu_id(cdu_map, r.get("cdu_name"))
        if not cdu_id:
            continue
        summary_date = r["snapshot_date"]
        key = (cdu_id, summary_date)
        existing = summaries_by_key.get(key)
        if existing is None:
            existing = db.execute(select(PortfolioSummary).where(
                PortfolioSummary.cdu_id == cdu_id,
                PortfolioSummary.summary_date == summary_date,
            )).scalars().first()
        kwargs = dict(
            total_mv_prev=r.get("total_mv_prev") or 0.0,
            total_daily_change=r.get("total_daily_change") or 0.0,
            total_mv_current=r.get("total_mv_current") or 0.0,
            ytm_weighted=r.get("ytm_weighted") or 0.0,
            duration_weighted=r.get("duration_weighted") or 0.0,
            benchmark_duration=r.get("benchmark_duration"),
            duration_status=r.get("duration_status"),
        )
        if existing and key in summaries_by_key:
            existing_values = {
                "ytm_weighted": existing.ytm_weighted,
                "duration_weighted": existing.duration_weighted,
            }
            if _summary_quality(existing_values) > _summary_quality(kwargs):
                continue
        if existing:
            for k, v in kwargs.items():
                setattr(existing, k, v)
            summary = existing
        else:
            summary = PortfolioSummary(cdu_id=cdu_id, summary_date=summary_date, **kwargs)
            db.add(summary)
        if key not in summaries_by_key:
            summary_count += 1
        summaries_by_key[key] = summary

    for (cdu_id, summary_date), summary in summaries_by_key.items():
        fund_totals_by_date[summary_date] = (
            fund_totals_by_date.get(summary_date, 0.0)
            + (summary.total_mv_current or 0.0)
        )
    for (cdu_id, summary_date), summary in summaries_by_key.items():
        fund_total = fund_totals_by_date.get(summary_date) or 0.0
        summary.cdu_share_pct = summary.total_mv_current / fund_total if fund_total else 0.0

    for r in position_rows:
        cdu_id = _resolve_cdu_id(cdu_map, r.get("cdu_name"))
        if not cdu_id:
            continue
        position_date = r["snapshot_date"]
        category = r["category"]
        existing_rows = db.execute(select(PortfolioPosition).where(
            PortfolioPosition.cdu_id == cdu_id,
            PortfolioPosition.position_date == position_date,
            PortfolioPosition.instrument_category == category,
            PortfolioPosition.instrument_code.is_(None),
        )).scalars().all()
        kwargs = dict(
            instrument_name=r.get("instrument_name"),
            nominal_volume=0.0,
            current_price=None,
            accrued_interest=0.0,
            market_value_current=r.get("market_value_current") or 0.0,
            market_value_prev=r.get("market_value_prev") or 0.0,
            daily_change=r.get("daily_change") or 0.0,
            pct_of_total=r.get("pct_of_total") or 0.0,
            ytm=r.get("ytm"),
            duration=r.get("duration"),
            hard_limit_status=r.get("hard_limit") or "ok",
            soft_limit_status=r.get("soft_limit") or "ok",
            free_limit_mln=r.get("free_limit_mln"),
            notes="rr_report_import",
        )
        if existing_rows:
            target = existing_rows[0]
            for extra in existing_rows[1:]:
                db.delete(extra)
            for k, v in kwargs.items():
                setattr(target, k, v)
        else:
            db.add(PortfolioPosition(
                cdu_id=cdu_id,
                position_date=position_date,
                instrument_code=None,
                instrument_category=category,
                **kwargs,
            ))
        position_count += 1

    return summary_count, position_count


def _resolve_cdu_id(cdu_map: dict[str, int], name: Optional[str]) -> Optional[int]:
    from services.cdu_lookup import resolve_name_id
    return resolve_name_id(cdu_map, name)


def _delete_rr_snapshots_for_dates(db: Session, model, date_column, dates: Iterable[date]) -> None:
    affected = sorted({d for d in dates if d is not None})
    for i in range(0, len(affected), 500):
        chunk = affected[i:i + 500]
        db.query(model).filter(
            date_column.in_(chunk),
            model.source == "rr_import",
        ).delete(synchronize_session=False)


def _delete_rr_detail_rows_for_dates(db: Session, dates: Iterable[date]) -> None:
    affected = sorted({d for d in dates if d is not None})
    if not affected:
        return
    rr_doc_ids = select(SourceDocument.id).where(SourceDocument.doc_type == "RISK_REPORT_XLSM")
    for model, date_column in (
        (BondLot, BondLot.valuation_date),
        (RepoLot, RepoLot.valuation_date),
        (DepositLot, DepositLot.valuation_date),
        # AccountReceivable is NOT deleted: it carries a transaction history
        # and is reconciled against the snapshot by ``_upsert_ar``.
    ):
        for i in range(0, len(affected), 500):
            chunk = affected[i:i + 500]
            db.query(model).filter(
                date_column.in_(chunk),
                model.source_doc_id.in_(rr_doc_ids),
            ).delete(synchronize_session=False)


def _upsert_cash(db: Session, rows: list[dict], cdu_map: dict[str, int],
                 source_doc_id: int) -> int:
    _delete_rr_snapshots_for_dates(
        db, CashSnapshot, CashSnapshot.snapshot_date,
        (r.get("snapshot_date") for r in rows),
    )
    cnt = 0
    for r in rows:
        cdu_id = _resolve_cdu_id(cdu_map, r["cdu_name"])
        if not cdu_id:
            continue
        existing = db.execute(select(CashSnapshot).where(
            CashSnapshot.cdu_id == cdu_id,
            CashSnapshot.snapshot_date == r["snapshot_date"],
            CashSnapshot.currency == r["currency"],
        )).scalars().first()
        if existing:
            existing.amount = r["amount"]
            existing.source = "rr_import"
            existing.source_doc_id = source_doc_id
        else:
            db.add(CashSnapshot(
                cdu_id=cdu_id,
                snapshot_date=r["snapshot_date"],
                currency=r["currency"],
                amount=r["amount"],
                source="rr_import",
                source_doc_id=source_doc_id,
            ))
        cnt += 1
    return cnt


def _upsert_mv(db: Session, rows: list[dict], cdu_map: dict[str, int],
               source_doc_id: int) -> int:
    _delete_rr_snapshots_for_dates(
        db, MVSnapshot, MVSnapshot.snapshot_date,
        (r.get("snapshot_date") for r in rows),
    )
    cnt = 0
    for r in rows:
        cdu_id = _resolve_cdu_id(cdu_map, r["cdu_name"])
        if not cdu_id:
            continue
        existing = db.execute(select(MVSnapshot).where(
            MVSnapshot.cdu_id == cdu_id,
            MVSnapshot.snapshot_date == r["snapshot_date"],
        )).scalars().first()
        kwargs = dict(
            cash_flow=r.get("cash_flow"),
            market_value_total=r["market_value_total"],
            return_pct=r.get("return_pct"),
            ytm_weighted=r.get("ytm_weighted"),
            duration_weighted=r.get("duration_weighted"),
            source="rr_import",
            source_doc_id=source_doc_id,
        )
        if existing:
            for k, v in kwargs.items():
                setattr(existing, k, v)
        else:
            db.add(MVSnapshot(
                cdu_id=cdu_id,
                snapshot_date=r["snapshot_date"],
                **kwargs,
            ))
        cnt += 1
    return cnt


def _upsert_reference(db: Session, rows: list[dict]) -> int:
    """Upsert Справочника выпусков по ISIN.

    Примечание по идемпотентности: в Risk Report один и тот же ISIN может
    встречаться несколько раз (разные строки НБРК/КФГД, исторические дубли).
    С ``autoflush=False`` SELECT не видит pending insert’ы — поэтому держим
    локальный кэш в рамках одного вызова.
    """
    pending: dict[str, InstrumentReference] = {}
    cnt = 0
    UPDATABLE = (
        "ticker_kase", "instrument_name", "security_type", "issuer",
        "issuer_type", "bond_type", "coupon_rate_pct", "coupon_rate_text",
        "coupon_type", "frequency", "base", "nominal",
        "start_date", "maturity_date", "last_coupon_date",
        "next_coupon_date", "quotation_type", "currency", "rating",
        "rating_source", "rating_agency",
    )
    for r in rows:
        isin = r.get("isin")
        if not isin:
            continue
        target = pending.get(isin)
        if target is None:
            target = db.execute(select(InstrumentReference).where(
                InstrumentReference.isin == isin,
            )).scalars().first()
        if target is not None:
            for k in UPDATABLE:
                v = r.get(k)
                if v is not None:
                    setattr(target, k, v)
        else:
            target = InstrumentReference(**{k: v for k, v in r.items() if k != "src_row"})
            db.add(target)
        pending[isin] = target
        cnt += 1
    return cnt


def _upsert_fx(db: Session, rows: list[dict]) -> int:
    pending: dict[tuple, FXRate] = {}
    cnt = 0
    for r in rows:
        key = (r["rate_date"], r["currency"])
        target = pending.get(key)
        if target is None:
            target = db.execute(select(FXRate).where(
                FXRate.rate_date == r["rate_date"],
                FXRate.currency == r["currency"],
            )).scalars().first()
        if target is not None:
            target.rate = r["rate"]
            target.source = "nbrk_rr"
        else:
            target = FXRate(
                rate_date=r["rate_date"],
                currency=r["currency"],
                rate=r["rate"],
                source="nbrk_rr",
            )
            db.add(target)
        pending[key] = target
        cnt += 1
    return cnt


def _upsert_mbm(db: Session, rows: list[dict]) -> int:
    pending: dict = {}
    cnt = 0
    for r in rows:
        key = r["index_date"]
        target = pending.get(key)
        if target is None:
            target = db.execute(select(MBMIndex).where(
                MBMIndex.index_date == key,
            )).scalars().first()
        if target is not None:
            if r.get("ytm_value") is not None:
                target.ytm_value = r["ytm_value"]
            if r.get("duration") is not None:
                target.duration = r["duration"]
            target.source = "rr_import"
        else:
            target = MBMIndex(
                index_date=key,
                ytm_value=r.get("ytm_value"),
                duration=r.get("duration"),
                source="rr_import",
            )
            db.add(target)
        pending[key] = target
        cnt += 1
    return cnt


def _upsert_bond_lots(db: Session, rows: list[dict], cdu_map: dict[str, int],
                      source_doc_id: int, snapshot_date: date) -> int:
    """Лоты ЦБ — снимок портфеля на snapshot_date («1 сделка = 1 лот»).

    Все прежние RR-лоты этой даты удалены ``_delete_rr_detail_rows_for_dates``,
    поэтому каждая строка листа вставляется как отдельный лот: две покупки
    одного ISIN остаются двумя лотами (раньше вторая могла перезаписать первую
    в зависимости от настроек autoflush сессии).
    """
    cnt = 0
    db.flush()  # make pending Справочник rows visible to the enrichment lookups
    ref_cache: dict[str, InstrumentReference] = {}
    for r in rows:
        cdu_id = _resolve_cdu_id(cdu_map, r["cdu_name"])
        if not cdu_id:
            continue
        val_date = r.get("valuation_date") or snapshot_date
        face = r.get("face_value") or 0.0
        qty = r.get("quantity") or 0.0
        db.add(BondLot(
            cdu_id=cdu_id, isin=r["isin"],
            category=r["category"],
            trade_date=r.get("trade_date") or val_date,
            valuation_date=val_date,
            quantity_initial=qty,
            quantity_current=qty,
            face_value_initial=face,
            face_value_current=face,
            purchase_price=r.get("purchase_price"),
            market_price=r.get("market_price"),
            accrued_interest=r.get("accrued_interest"),
            market_value=r.get("market_value"),
            total_value=r.get("total_value"),
            weight=r.get("weight"),
            ytm=r.get("ytm"),
            duration=r.get("duration"),
            maturity_status=r.get("maturity_status"),
            source_doc_id=source_doc_id,
        ))
        _enrich_reference_from_lot(db, r, ref_cache)
        cnt += 1
    db.flush()
    return cnt


def _enrich_reference_from_lot(db: Session, r: dict, cache: dict) -> None:
    """Fill gaps of the Справочник from the lot sheet (maturity, coupon, base).

    Valuation on dates without a Risk Report needs the coupon schedule; the
    instrument sheets always carry it even when the «Справочник» row is
    incomplete.  Existing values are never overwritten here.
    """
    isin = (r.get("isin") or "").strip().upper()
    if not isin:
        return
    ref = cache.get(isin)
    if ref is None:
        ref = db.execute(select(InstrumentReference).where(
            InstrumentReference.isin == isin,
        )).scalars().first()
    if ref is None:
        ref = InstrumentReference(isin=isin, currency="USD" if r.get("category") == "FOREIGN_BONDS" else "KZT",
                                  source="RISK_REPORT_LOT")
        db.add(ref)
    cache[isin] = ref
    coupon = r.get("coupon_rate")
    if coupon is not None and abs(coupon) <= 1:
        coupon = coupon * 100.0
    category_label = {
        "GOV_BONDS": "ГЦБ", "AGENCY_BONDS": "Агентские",
        "MFO_BONDS": "МФО", "FOREIGN_BONDS": "Ин. ЦБ",
    }.get(r.get("category"))
    for field, value in (
        ("maturity_date", r.get("maturity_date")),
        ("coupon_rate_pct", coupon),
        ("frequency", r.get("frequency")),
        ("base", r.get("base")),
        ("issuer", r.get("issuer")),
        ("issuer_type", category_label),
        ("bond_type", category_label),
    ):
        if value is not None and getattr(ref, field, None) in (None, ""):
            setattr(ref, field, value)


def _upsert_repo_lots(db: Session, rows: list[dict], cdu_map: dict[str, int],
                      source_doc_id: int, snapshot_date: date) -> int:
    cnt = 0
    for r in rows:
        cdu_id = _resolve_cdu_id(cdu_map, r["cdu_name"])
        if not cdu_id:
            continue
        val_date = r.get("valuation_date") or snapshot_date
        db.add(RepoLot(
            cdu_id=cdu_id,
            deal_id=r.get("deal_id"),
            instrument_code=r.get("instrument_code"),
            isin=r.get("isin"),
            trade_date=r.get("trade_date") or val_date,
            valuation_date=val_date,
            close_date=r.get("close_date"),
            face_value=r["face_value"],
            close_value=r.get("close_value"),
            repo_rate_pct=r.get("repo_rate_pct"),
            base=r.get("base") or 366,
            accrued_interest=r.get("accrued_interest"),
            term_days=r.get("term_days"),
            market_value=r.get("market_value"),
            ytm=r.get("ytm"),
            duration=r.get("duration"),
            maturity_status=r.get("maturity_status"),
            is_closed=bool(r.get("close_date") and r["close_date"] <= val_date),
            source_doc_id=source_doc_id,
        ))
        cnt += 1
    db.flush()
    return cnt


def _upsert_deposit_lots(db: Session, rows: list[dict], cdu_map: dict[str, int],
                         source_doc_id: int, snapshot_date: date) -> int:
    cnt = 0
    for r in rows:
        cdu_id = _resolve_cdu_id(cdu_map, r["cdu_name"])
        if not cdu_id:
            continue
        val_date = r.get("valuation_date") or snapshot_date
        db.add(DepositLot(
            cdu_id=cdu_id,
            trade_date=r.get("trade_date") or val_date,
            valuation_date=val_date,
            close_date=r.get("close_date"),
            principal=r["principal"],
            interest_rate_pct=r.get("interest_rate_pct"),
            accrued_interest=r.get("accrued_interest"),
            market_value=r.get("market_value"),
            source_doc_id=source_doc_id,
        ))
        cnt += 1
    db.flush()
    return cnt


def _upsert_ar(db: Session, rows: list[dict], cdu_map: dict[str, int],
               source_doc_id: int, snapshot_date: Optional[date] = None) -> int:
    """Дебиторская задолженность из листа «Accounts receivable».

    Лист — это *снимок* открытых остатков на дату отчёта.  Поэтому:
    * строка с тем же (ЧДУ, ISIN, дата постановки) обновляет существующую
      запись (в т.ч. созданную движком купонов/погашений — без дублей);
    * изменение остатка фиксируется операцией ReceivableTransaction на дату
      снимка (история «на дату» сохраняется);
    * открытая RR-дебиторка, которой больше нет в снимке (или остаток 0),
      закрывается операцией REPAYMENT — раньше она оставалась OPEN навсегда и
      завышала портфель.
    """
    from models.db_models import ReceivableTransaction
    from api.routes.accounting import receivable_balance_as_of  # local: avoid import cycle

    cnt = 0
    seen_ids: set[int] = set()
    present_cdus: set[int] = set()
    for r in rows:
        cdu_id = _resolve_cdu_id(cdu_map, r["cdu_name"])
        if not cdu_id:
            # Для НБ РК вкладок AR может не содержать ДУ — пропускаем строки без cdu
            continue
        present_cdus.add(cdu_id)
        rec_date = r.get("record_date") or snapshot_date
        if not rec_date:
            continue
        balance = float(r.get("balance_kzt") or 0.0)
        existing = db.execute(select(AccountReceivable).where(
            AccountReceivable.cdu_id == cdu_id,
            AccountReceivable.isin == r["isin"],
            AccountReceivable.record_date == rec_date,
            AccountReceivable.is_manual.is_(False),
        )).scalars().first()
        if existing is None:
            existing = AccountReceivable(
                cdu_id=cdu_id,
                record_date=rec_date,
                isin=r["isin"],
                description=r.get("description"),
                currency=r.get("currency", "KZT"),
                amount=balance,
                amount_kzt=balance,
                balance_currency=r.get("balance_currency"),
                balance_kzt=balance,
                due_date=r.get("due_date"),
                status="OPEN",
                related_event_type=_ar_event_type(r.get("description")),
                workflow_status="APPROVED",
                is_manual=False,
                is_active=True,
                source_doc_id=source_doc_id,
                created_by="RISK_REPORT",
            )
            db.add(existing)
            db.flush()
        else:
            existing.description = r.get("description") or existing.description
            existing.due_date = r.get("due_date") or existing.due_date
            if snapshot_date:
                current = receivable_balance_as_of(db, existing, snapshot_date)
                delta = balance - current
                if abs(delta) > 0.005:
                    db.add(ReceivableTransaction(
                        receivable_id=existing.id,
                        operation_date=snapshot_date,
                        operation_type="ADJUSTMENT" if delta > 0 else "REPAYMENT",
                        amount=abs(delta) if delta < 0 else delta,
                        amount_kzt=abs(delta) if delta < 0 else delta,
                        currency=existing.currency or "KZT",
                        basis=f"Остаток по Risk Report за {snapshot_date:%d.%m.%Y}",
                        workflow_status="APPROVED", is_active=True,
                        created_by="RISK_REPORT", approved_by="SYSTEM",
                    ))
            existing.balance_kzt = balance
            existing.balance_currency = r.get("balance_currency")
            existing.status = "OPEN" if balance > 0.005 else "CLOSED"
            existing.source_doc_id = source_doc_id
        seen_ids.add(existing.id)
        cnt += 1

    if snapshot_date is not None and present_cdus:
        rr_doc_ids = select(SourceDocument.id).where(SourceDocument.doc_type == "RISK_REPORT_XLSM")
        stale = db.execute(select(AccountReceivable).where(
            AccountReceivable.cdu_id.in_(present_cdus),
            AccountReceivable.is_manual.is_(False),
            AccountReceivable.is_active.is_(True),
            AccountReceivable.record_date <= snapshot_date,
            AccountReceivable.status.in_(("OPEN", "PARTIAL")),
            AccountReceivable.source_doc_id.in_(rr_doc_ids),
        )).scalars().all()
        for ar in stale:
            if ar.id in seen_ids:
                continue
            current = receivable_balance_as_of(db, ar, snapshot_date)
            if current > 0.005:
                db.add(ReceivableTransaction(
                    receivable_id=ar.id,
                    operation_date=snapshot_date,
                    operation_type="REPAYMENT",
                    amount=current, amount_kzt=current,
                    currency=ar.currency or "KZT",
                    basis=f"Отсутствует в Risk Report за {snapshot_date:%d.%m.%Y} — погашена",
                    workflow_status="APPROVED", is_active=True,
                    created_by="RISK_REPORT", approved_by="SYSTEM",
                ))
            ar.status = "CLOSED"
            ar.balance_kzt = 0.0
            ar.actual_value_date = ar.actual_value_date or snapshot_date
    db.flush()
    return cnt


def _ar_event_type(description: Optional[str]) -> Optional[str]:
    text = (description or "").lower()
    if "купон" in text or "вознагражд" in text:
        return "COUPON"
    if "погашен" in text:
        return "REDEMPTION"
    return None
