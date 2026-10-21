"""KASE «TONIA Compounded»: TCI and TCR_1M / TCR_3M / TCR_6M by calculation date.

Source: the Excel file of the KASE page «TONIA Compounded»
(https://kase.kz/ru/indexes-and-indicators/repo/tonia-compounded → «Скачать
архивные данные»), for one date or a period::

    TONIA Compounded 2026-10-14 - 2026-10-14
    Дата расчетов | TCR_1M  | TCR_3M  | TCR_6M  | TCI
    14.10.2026    | 16.1091 | 16.4560 | 17.4990 | 2.2739959262

Values are stored per indicator and date; a repeated upload of a date
updates the values (no duplicates).

Floating coupons: rate of a period = TCR of the issue (Справочник ЦБ →
«Индикатор плавающей ставки») on the fixing date + margin of the issue.  When
there is no value on the fixing date the last value published before it is
used.
"""
from __future__ import annotations

import io
import hashlib
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import httpx
from loguru import logger
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from config import settings
from models.db_models import SourceDocument, ToniaCompounded
from services.audit import write_audit
from services.parser.number_utils import parse_kz_date, parse_kz_number

INDICATORS = ("TCI", "TCR_1M", "TCR_3M", "TCR_6M")
RATE_INDICATORS = ("TCR_1M", "TCR_3M", "TCR_6M")
FILE_RE = re.compile(r"(?i)tonia\s*_?compounded")


class ToniaClient:
    """Download the official TONIA Compounded archive from KASE."""

    def __init__(self, *, url: Optional[str] = None,
                 transport: Optional[httpx.AsyncBaseTransport] = None) -> None:
        self.url = url or settings.kase_tonia_xlsx_url
        self.transport = transport
        self.headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
            ),
            "Accept-Language": "ru,en;q=0.8",
            "Referer": "https://kase.kz/ru/indexes-and-indicators/repo/tonia-compounded",
        }

    async def fetch_archive(self, start: date, end: date) -> bytes:
        if start > end:
            raise ValueError("Начальная дата архива TONIA позже конечной даты")
        params = {
            "start_date": start.isoformat(),
            "end_date": end.isoformat(),
            "language": "ru",
        }
        async with httpx.AsyncClient(
            timeout=60.0,
            headers=self.headers,
            follow_redirects=True,
            transport=self.transport,
        ) as client:
            response = await client.get(self.url, params=params)
            response.raise_for_status()
        content = response.content
        if not content.startswith(b"PK"):
            content_type = response.headers.get("Content-Type", "")
            raise ValueError(
                f"KASE вернула архив TONIA не в формате XLSX: тип {content_type or 'не указан'}, размер {len(content)} байт"
            )
        parsed = parse_tonia_xlsx(content)
        if not parsed["rows"]:
            raise ValueError("В архиве TONIA с KASE нет значений показателей")
        logger.info(
            f"KASE TONIA archive {start.isoformat()}..{end.isoformat()}: "
            f"{len(parsed['rows'])} dates, {len(content)} bytes"
        )
        return content


def looks_like_tonia_file(name: str) -> bool:
    return bool(FILE_RE.search(name or ""))


def _norm(v: Any) -> str:
    return re.sub(r"[\s\-]+", "_", str(v or "").strip().upper())


def parse_tonia_xlsx(source: bytes | str | Path) -> Dict[str, Any]:
    """Rows ``{"date": date, "TCI": float, "TCR_1M": ...}`` and warnings."""
    data = source if isinstance(source, bytes) else Path(source).read_bytes()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header_idx = None
    cols: Dict[str, int] = {}
    for i, row in enumerate(rows[:20]):
        names = [_norm(v) for v in row]
        if any(n.startswith("ДАТА") for n in names) and any(n in INDICATORS for n in names):
            header_idx = i
            for j, n in enumerate(names):
                if n.startswith("ДАТА"):
                    cols.setdefault("date", j)
                elif n in INDICATORS:
                    cols[n] = j
            break
    if header_idx is None:
        raise ValueError("Не найдена шапка «Дата расчетов | TCR_1M | TCR_3M | TCR_6M | TCI» — "
                         "загрузите файл со страницы KASE «TONIA Compounded»")
    out: List[Dict[str, Any]] = []
    warnings: List[str] = []
    for n, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not row or all(v in (None, "") for v in row):
            continue
        raw_date = row[cols["date"]] if cols["date"] < len(row) else None
        d = raw_date.date() if isinstance(raw_date, datetime) else (
            raw_date if isinstance(raw_date, date) else parse_kz_date(raw_date))
        if d is None:
            warnings.append(f"Строка {n}: не распознана дата «{raw_date}» — пропущена")
            continue
        item: Dict[str, Any] = {"date": d}
        for ind in INDICATORS:
            j = cols.get(ind)
            if j is None or j >= len(row):
                continue
            v = row[j]
            value = float(v) if isinstance(v, (int, float)) else parse_kz_number(v)
            if value is not None:
                item[ind] = float(value)
        if len(item) > 1:
            out.append(item)
    missing = [i for i in INDICATORS if i not in cols]
    if missing:
        warnings.append("В файле нет колонок: " + ", ".join(missing))
    return {"rows": out, "warnings": warnings}


def upsert(db: Session, rows: List[Dict[str, Any]], *, source_doc_id: Optional[int] = None,
           actor: Optional[str] = None, source: str = "kase_xlsx") -> Dict[str, int]:
    inserted = updated = unchanged = 0
    for item in rows:
        for ind in INDICATORS:
            if ind not in item:
                continue
            row = db.execute(select(ToniaCompounded).where(
                ToniaCompounded.calc_date == item["date"], ToniaCompounded.indicator == ind,
            )).scalars().first()
            if row is None:
                db.add(ToniaCompounded(calc_date=item["date"], indicator=ind, value=item[ind],
                                       source=source, source_doc_id=source_doc_id, updated_by=actor))
                inserted += 1
            else:
                if abs((row.value or 0.0) - item[ind]) > 1e-12:
                    updated += 1
                    row.value = item[ind]
                    row.source = source
                    row.source_doc_id = source_doc_id
                    row.updated_by = actor
                else:
                    unchanged += 1
    db.flush()
    dates = sorted({r["date"] for r in rows})
    return {"dates": len(dates), "inserted": inserted, "updated": updated, "unchanged": unchanged,
            "date_from": dates[0] if dates else None, "date_to": dates[-1] if dates else None}


def import_archive(
    db: Session,
    raw: bytes,
    *,
    filename: str,
    actor: str,
    source: str = "kase_api",
) -> Dict[str, Any]:
    """Parse, archive and idempotently store one KASE TONIA workbook."""
    parsed = parse_tonia_xlsx(raw)
    rows = parsed["rows"]
    if not rows:
        raise ValueError("В файле KASE нет значений TONIA Compounded")

    digest = hashlib.sha256(raw).hexdigest()
    doc = db.execute(select(SourceDocument).where(
        SourceDocument.doc_type == "TONIA_COMPOUNDED",
        SourceDocument.sha256 == digest,
    )).scalars().first()
    duplicate_document = doc is not None
    if doc is None:
        folder = Path(settings.upload_dir) / "tonia"
        folder.mkdir(parents=True, exist_ok=True)
        safe_name = Path(filename).name or "tonia_compounded.xlsx"
        path = folder / f"{datetime.utcnow():%Y%m%d_%H%M%S}_{digest[:12]}_{safe_name}"
        path.write_bytes(raw)
        doc = SourceDocument(
            doc_type="TONIA_COMPOUNDED",
            doc_date=max(row["date"] for row in rows),
            file_name=safe_name,
            file_path=str(path),
            sha256=digest,
            file_size=len(raw),
            uploaded_by=actor,
            parsed_at=datetime.utcnow(),
            parse_status="OK",
            rows_imported=len(rows),
        )
        db.add(doc)
        db.flush()

    counters = upsert(db, rows, source_doc_id=doc.id, actor=actor, source=source)
    details = {
        key: value.isoformat() if isinstance(value, date) else value
        for key, value in counters.items()
    }
    details.update({"sha256": digest, "duplicate_document": duplicate_document, "source": source})
    write_audit(
        db,
        user=actor,
        action="TONIA_AUTO_IMPORTED" if source == "kase_api" else "TONIA_IMPORTED",
        entity="ToniaCompounded",
        entity_id=doc.id,
        details=details,
    )
    return {
        **counters,
        "warnings": parsed["warnings"],
        "source_doc_id": doc.id,
        "duplicate_document": duplicate_document,
    }


async def refresh_from_kase(
    db: Session,
    *,
    actor: str = "scheduler",
    client: Optional[ToniaClient] = None,
    today: Optional[date] = None,
) -> Dict[str, Any]:
    """Fetch recent KASE values and import them; overlap catches revisions."""
    end = today or date.today()
    latest = db.execute(select(func.max(ToniaCompounded.calc_date))).scalar()
    if latest is None:
        start = settings.kase_tonia_start_date
    else:
        lookback = max(0, int(settings.kase_tonia_lookback_days))
        start = max(settings.kase_tonia_start_date, latest - timedelta(days=lookback))
    raw = await (client or ToniaClient()).fetch_archive(start, end)
    filename = f"tonia_compounded_{start.isoformat()}_{end.isoformat()}.xlsx"
    result = import_archive(db, raw, filename=filename, actor=actor, source="kase_api")
    return {**result, "requested_from": start, "requested_to": end}


def import_file(db: Session, path: str | Path, *, source_doc_id: Optional[int] = None,
                actor: Optional[str] = None) -> Dict[str, Any]:
    parsed = parse_tonia_xlsx(Path(path))
    counters = upsert(db, parsed["rows"], source_doc_id=source_doc_id, actor=actor)
    return {**counters, "warnings": parsed["warnings"]}


# ───────────────────────── lookup ─────────────────────────

class ToniaBook:
    """Cached lookup «last value on or before a date» per indicator."""

    def __init__(self, db: Session):
        self.db = db
        self._series: Optional[Dict[str, List[Tuple[date, float]]]] = None

    def _load(self) -> Dict[str, List[Tuple[date, float]]]:
        if self._series is None:
            series: Dict[str, List[Tuple[date, float]]] = {}
            try:
                rows = self.db.execute(select(ToniaCompounded).order_by(ToniaCompounded.calc_date.asc())).scalars()
                for r in rows:
                    series.setdefault(r.indicator.upper(), []).append((r.calc_date, float(r.value)))
            except Exception:  # noqa: BLE001 — table missing in an old database
                series = {}
            self._series = series
        return self._series

    def value_on(self, indicator: str, d: date) -> Optional[Tuple[date, float]]:
        """(date of the value, value in %) — the last value published on or before ``d``."""
        series = self._load().get((indicator or "").upper())
        if not series:
            return None
        lo, hi = 0, len(series)
        while lo < hi:
            mid = (lo + hi) // 2
            if series[mid][0] <= d:
                lo = mid + 1
            else:
                hi = mid
        return series[lo - 1] if lo else None


def float_rate(book: ToniaBook, index: Optional[str], margin_pct: Optional[float],
               period_start: date, lag_days: int = 0) -> Optional[Tuple[float, date, float]]:
    """Annual rate (fraction) of a floating coupon period: TCR + margin.

    Returns (rate, date of the TCR value, TCR %) or None when the index is
    unknown or has no value before the fixing date.
    """
    if not index:
        return None
    found = book.value_on(index, period_start - timedelta(days=int(lag_days or 0)))
    if found is None:
        return None
    on, tcr = found
    return (tcr + float(margin_pct or 0.0)) / 100.0, on, tcr


def detect_index(text: Optional[str]) -> Tuple[Optional[str], Optional[float]]:
    """«TCR_3M + 1,5%» → ("TCR_3M", 1.5) from the coupon text of the Справочник."""
    if not text:
        return None, None
    t = text.upper().replace(" ", "")
    m = re.search(r"TCR_?(1|3|6)M", t)
    if not m:
        return None, None
    index = f"TCR_{m.group(1)}M"
    mm = re.search(r"TCR_?\dM([+\-])(\d+(?:[.,]\d+)?)", t)
    margin = None
    if mm:
        margin = float(mm.group(2).replace(",", "."))
        if mm.group(1) == "-":
            margin = -margin
    return index, margin
