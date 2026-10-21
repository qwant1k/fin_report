"""Regression tests for the testers' requirement log of 15.10.2026.

Files of the log (Trade Reports Halyk 14.09 and Tansar 30.09, TONIA Compounded)
are read from ``Примеры/Журнал 15.10.2026`` when available.  The Risk Report
of 11.06.2026 is password-protected, so its situations are reproduced with the
same figures directly in the database.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from models.db_models import (
    AccountReceivable, BondLot, CashMovement, CashSnapshot, CDU, CouponSchedule, InstrumentReference,
    PortfolioPosition, RepoLot, SourceDocument, ToniaCompounded, Trade,
)
from tests.test_requirements_20260929 import LOG_DIR as LOG_0929, _cdu, _upload, stand  # noqa: F401

LOG = LOG_0929.parent / "Журнал 15.10.2026"
HALYK_1409 = LOG / "Trade report 1409 2026.xlsx"
TANSAR_3009 = LOG / "Приложение 1 Trade report 30.09.2026 Тансар.xlsx"
TONIA = LOG / "TONIACompounded2026-10-14-2026-10-14.xlsx"
RR_3112 = LOG_0929.parent / "risk report_31122025_.xlsm"


def _rr_doc(db, on: date) -> SourceDocument:
    doc = SourceDocument(doc_type="RISK_REPORT_XLSM", doc_date=on, file_name=f"risk report_{on:%d%m%Y}_.xlsm",
                         file_path="x", parse_status="OK")
    db.add(doc)
    db.flush()
    return doc


def _cash(client, cdu_id: int, on: str) -> dict:
    r = client.get("/api/dashboard/category-breakdown", params={"cdu_id": cdu_id, "category": "CASH", "date": on})
    assert r.status_code == 200, r.text
    return r.json()


# ───────────── п.2: deals of REPO from the Risk Report move the cash ─────────────

CLOSES_1206 = [78_046_623.05, 502_516_364.94, 3_293_208_277.16, 154_214_741.10,
               70_096_144.01, 1_809_107_323.79, 37_020_661.42]          # 7 closings = 5 944 210 135,47
OPENS_1206 = [2_000_000_000.00, 1_500_000_000.00, 1_444_348_448.70, 600_000_000.00, 400_000_000.00]


def test_repo_closes_of_the_slice_and_new_openings_move_the_cash(stand):
    """Halyk 12.06.2026: 1 246 977,50 − 5 944 348 448,70 + 5 944 210 135,47 = 1 108 664,27.

    The Risk Report of 11.06 carries its cash forward to 12.06 with a zero cash
    flow; that row must not hide the movements of 12.06.
    """
    client, Session, _ = stand
    with Session() as db:
        halyk = _cdu(Session, "Halyk Finance")
        doc = _rr_doc(db, date(2026, 6, 11))
        db.add(BondLot(cdu_id=halyk.id, isin="KZKD00000535", instrument_code="MUM180_0006", category="GOV_BONDS",
                       trade_date=date(2024, 7, 18), valuation_date=date(2026, 6, 11), face_value_current=0.0,
                       source_doc_id=doc.id))
        for d in (date(2026, 6, 11), date(2026, 6, 12)):          # 12.06 = carried forward row
            db.add(CashSnapshot(cdu_id=halyk.id, snapshot_date=d, currency="KZT", amount=1_246_977.50,
                                source="rr_import", source_doc_id=doc.id))
        for i, close in enumerate(CLOSES_1206):
            db.add(RepoLot(cdu_id=halyk.id, deal_id=f"1232162{i:02d}", trade_date=date(2026, 6, 11),
                           valuation_date=date(2026, 6, 11), close_date=date(2026, 6, 12),
                           face_value=close - 1000.0, close_value=close, base=366, source_doc_id=doc.id))
        tr = SourceDocument(doc_type="TRADE_REPORT", doc_date=date(2026, 6, 12), file_name="tr 12.06", file_path="x")
        db.add(tr)
        db.flush()
        for i, amount in enumerate(OPENS_1206):
            db.add(CashMovement(cdu_id=halyk.id, operation_date=date(2026, 6, 12), amount=-amount,
                                amount_kzt=-amount, currency="KZT", movement_type="REPO_OPEN",
                                basis=f"Автоматически по сделке 1232300{i}", is_manual=False,
                                workflow_status="APPROVED", source_doc_id=tr.id))
        db.commit()
        halyk_id = halyk.id
    cash = _cash(client, halyk_id, "2026-06-12")["cash"]
    assert cash["opening"]["date"] == "2026-06-11"
    assert sum(i["amount_kzt"] for i in cash["implied"] if i["counted"]) == pytest.approx(5_944_210_135.47)
    assert sum(m["amount_kzt"] for m in cash["movements"]) == pytest.approx(-5_944_348_448.70)
    assert cash["total"] == pytest.approx(1_108_664.27)


# ───────────── п.3 и п.4: Trade Reports Tansar 30.09 and Halyk 14.09 ─────────────

@pytest.mark.skipif(not TANSAR_3009.exists(), reason="файл Tansar 30.09.2026 недоступен")
def test_tansar_purchase_with_coded_kp_column(stand):
    client, Session, _ = stand
    body = _upload(client, TANSAR_3009)
    assert not any("к/п" in w.lower() for w in body.get("warnings", [])), body
    tansar = _cdu(Session, "Tansar Capital")
    with Session() as db:
        buy = db.query(Trade).filter_by(cdu_id=tansar.id, operation_type="BUY").one()
        assert (buy.instrument_code, buy.trade_date, buy.value_date) == ("MOM060_0054", date(2026, 9, 30), date(2026, 10, 2))
        assert buy.face_value == pytest.approx(1_900_000_000)
        mv = db.query(CashMovement).filter_by(related_trade_id=buy.id).one()
        assert (mv.operation_date, mv.amount_kzt) == (date(2026, 10, 2), pytest.approx(-1_998_969_311.11))
        assert db.query(Trade).filter_by(cdu_id=tansar.id, operation_type="REPO_OPEN").count() == 3
    from services.portfolio.engine import build_portfolio
    with Session() as db:
        assert not [b for b in build_portfolio(db, tansar.id, date(2026, 10, 1)).bonds if b.ticker == "MOM060_0054"]
        held = [b for b in build_portfolio(db, tansar.id, date(2026, 10, 2)).bonds if b.ticker == "MOM060_0054"]
        assert held and held[0].face_value == pytest.approx(1_900_000_000)


@pytest.mark.skipif(not HALYK_1409.exists(), reason="файл Halyk 14.09.2026 недоступен")
def test_halyk_sales_are_received_on_settlement_date(stand):
    client, Session, _ = stand
    _upload(client, HALYK_1409)
    halyk = _cdu(Session, "Halyk Finance")
    cash = _cash(client, halyk.id, "2026-09-16")["cash"]
    sales = sorted(round(m["amount_kzt"], 2) for m in cash["movements"]
                   if m["date"] == "2026-09-16" and m["amount_kzt"] > 0)
    assert sales == [3_246_026.18, 3_661_279_876.34]


# ───────────── п.5: ISIN column ─────────────

def test_breakdown_items_carry_isin_from_reference(stand):
    client, Session, _ = stand
    with Session() as db:
        halyk = _cdu(Session, "Halyk Finance")
        doc = _rr_doc(db, date(2026, 6, 11))
        db.add(InstrumentReference(isin="KZKD00001210", ticker_kase="MUM072_0014", nominal=1000,
                                   maturity_date=date(2030, 1, 1), coupon_rate_pct=10, frequency=1,
                                   start_date=date(2024, 1, 1), base="30/360"))
        db.add(BondLot(cdu_id=halyk.id, isin="MUM072_0014", instrument_code="MUM072_0014", category="GOV_BONDS",
                       trade_date=date(2024, 7, 18), valuation_date=date(2026, 6, 11),
                       face_value_current=1_000_000.0, market_price=95.0, accrued_interest=1000.0,
                       total_value=951_000.0, source_doc_id=doc.id))
        db.add(RepoLot(cdu_id=halyk.id, deal_id="123216296", instrument_code="MUM072_0014", isin="KZKD00001210",
                       trade_date=date(2026, 6, 11), valuation_date=date(2026, 6, 11), close_date=date(2026, 6, 15),
                       face_value=154_142_104.27, close_value=154_214_741.10, base=366, source_doc_id=doc.id))
        db.commit()
        halyk_id = halyk.id
    bonds = client.get("/api/dashboard/category-breakdown",
                       params={"cdu_id": halyk_id, "category": "GOV_BONDS", "date": "2026-06-11"}).json()
    assert bonds["items"][0]["isin"] == "KZKD00001210"
    repos = client.get("/api/dashboard/category-breakdown",
                       params={"cdu_id": halyk_id, "category": "REVERSE_REPO", "date": "2026-06-12"}).json()
    assert repos["items"][0]["isin"] == "KZKD00001210"
    assert repos["items"][0]["collateral_ticker"] == "MUM072_0014"
    assert repos["items"][0]["rate"] == pytest.approx((154_214_741.10 / 154_142_104.27 - 1) * 366 / 4)


# ───────────── п.6 и п.8: redemption and coupon events of past dates ─────────────

def _held_bond(db, cdu: CDU, *, isin, ticker, face, maturity, coupon=None, freq=None, start=None):
    doc = _rr_doc(db, date(2026, 6, 11))
    db.add(InstrumentReference(isin=isin, ticker_kase=ticker, nominal=100 if coupon is None else 1000,
                               maturity_date=maturity, coupon_rate_pct=coupon, frequency=freq,
                               start_date=start, base="30/360", bond_type="ГЦБ",
                               coupon_type="zero" if coupon is None else "fix",
                               issuer="Министерство финансов Республики Казахстан"))
    db.add(BondLot(cdu_id=cdu.id, isin=isin, instrument_code=ticker, category="GOV_BONDS",
                   trade_date=date(2026, 3, 18), valuation_date=date(2026, 6, 11), face_value_current=face,
                   market_price=99.0, accrued_interest=0.0, total_value=face * 0.99, source_doc_id=doc.id))
    db.add(CashSnapshot(cdu_id=cdu.id, snapshot_date=date(2026, 6, 11), currency="KZT", amount=1_000_000.0,
                        source="rr_import", source_doc_id=doc.id))


def test_matured_note_is_redeemed_in_cash_on_maturity(stand):
    client, Session, _ = stand
    halyk = _cdu(Session, "Halyk Finance")
    with Session() as db:
        _held_bond(db, halyk, isin="KZW100014429", ticker="NTK091_2876", face=9_000_000_000.0,
                   maturity=date(2026, 6, 17))
        db.commit()
    r = client.post("/api/calculate/", json={"report_date": "2026-06-17", "recalculate": True})
    assert r.status_code == 200, r.text
    with Session() as db:
        rows = {p.instrument_category: p.market_value_current for p in db.query(PortfolioPosition).filter_by(
            cdu_id=halyk.id, position_date=date(2026, 6, 17), instrument_code=None)}
    assert rows["GOV_BONDS"] == pytest.approx(0.0)
    assert rows["CASH"] == pytest.approx(9_001_000_000.0)
    cash = _cash(client, halyk.id, "2026-06-17")["cash"]
    assert [(m["date"], m["amount_kzt"]) for m in cash["movements"]] == [("2026-06-17", pytest.approx(9e9))]
    # on the day before the note is still a security
    client.post("/api/calculate/", json={"report_date": "2026-06-16", "recalculate": True})
    with Session() as db:
        gov = db.query(PortfolioPosition).filter_by(cdu_id=halyk.id, position_date=date(2026, 6, 16),
                                                    instrument_code=None, instrument_category="GOV_BONDS").one()
        assert gov.market_value_current > 8.9e9


def test_coupon_becomes_receivable_then_cash_on_actual_receipt(stand):
    client, Session, current = stand
    tansar = _cdu(Session, "Tansar Capital")
    with Session() as db:
        _held_bond(db, tansar, isin="KZKD00001160", ticker="MUM096_0012", face=3_750_000_000.0,
                   maturity=date(2029, 7, 28), coupon=10.55, freq=1, start=date(2021, 7, 28))
        db.commit()
    client.post("/api/calculate/", json={"report_date": "2026-07-28", "recalculate": True})
    with Session() as db:
        ar = db.query(AccountReceivable).filter_by(cdu_id=tansar.id, isin="KZKD00001160").one()
        assert (ar.record_date, ar.due_date, ar.status) == (date(2026, 7, 28), date(2026, 7, 28), "OPEN")
        assert ar.amount_kzt == pytest.approx(3_750_000_000 * 0.1055)
        assert "MUM096_0012" in ar.description
        ar_id = ar.id
        # no cash on the planned date
        assert not db.query(CashMovement).filter(CashMovement.cdu_id == tansar.id,
                                                 CashMovement.is_active.is_(True)).count()
    rec = client.get("/api/dashboard/category-breakdown",
                     params={"cdu_id": tansar.id, "category": "RECEIVABLES", "date": "2026-07-28"}).json()
    assert rec["total"] == pytest.approx(395_625_000)

    # the money came on 30.07: the user states the actual date
    r = client.post("/api/accounting/operations", json={
        "purpose": "RECEIVABLE_CLOSE", "cdu_id": tansar.id, "receivable_id": ar_id,
        "operation_date": "2026-07-30", "amount": "395 625 000",
    })
    assert r.status_code == 200, r.text
    tx = r.json()["created"][0]
    current["user"] = "checker"
    assert client.post(f"/api/accounting/{tx['entity']}/{tx['id']}/approve", json={}).status_code == 200
    current["user"] = "maker"
    client.post("/api/calculate/", json={"report_date": "2026-07-30", "recalculate": True})
    cash = _cash(client, tansar.id, "2026-07-30")["cash"]
    assert [(m["date"], round(m["amount_kzt"], 2)) for m in cash["movements"]] == [("2026-07-30", 395_625_000.0)]
    with Session() as db:
        rows = {p.instrument_category: p.market_value_current for p in db.query(PortfolioPosition).filter_by(
            cdu_id=tansar.id, position_date=date(2026, 7, 30), instrument_code=None)}
        assert rows["RECEIVABLES"] == pytest.approx(0.0, abs=0.01)
        assert rows["CASH"] == pytest.approx(1_000_000 + 395_625_000)


def test_coupon_settlement_parameter_books_cash_on_date(stand):
    client, Session, _ = stand
    tansar = _cdu(Session, "Tansar Capital")
    with Session() as db:
        _held_bond(db, tansar, isin="KZKD00001160", ticker="MUM096_0012", face=3_750_000_000.0,
                   maturity=date(2029, 7, 28), coupon=10.55, freq=1, start=date(2021, 7, 28))
        db.commit()
    assert client.put("/api/formulas/parameters/coupon_settlement",
                      json={"value": "cash", "reason": "купоны ГЦБ приходят день в день"}).status_code == 200
    try:
        client.post("/api/calculate/", json={"report_date": "2026-07-28", "recalculate": True})
        with Session() as db:
            assert not db.query(AccountReceivable).filter_by(cdu_id=tansar.id, is_active=True).count()
            mv = db.query(CashMovement).filter_by(cdu_id=tansar.id, is_active=True).one()
            assert (mv.operation_date, mv.movement_type) == (date(2026, 7, 28), "COUPON")
    finally:
        client.post("/api/formulas/parameters/coupon_settlement/reset", json={})


# ───────────── п.7: NBRK portfolios in «Сводка» ─────────────

@pytest.mark.skipif(not (LOG_0929 / "КФГД_20260925.csv").exists(), reason="файлы НБ РК недоступны")
def test_nbrk_file_gives_portfolio_cash_and_deposits(stand):
    client, Session, _ = stand
    for n in ("КФГД_20260924.csv", "КФГД_20260925.csv"):
        r = client.post("/api/primary-data/upload", files={"file": (n, (LOG_0929 / n).read_bytes())})
        assert r.status_code == 200, r.text
    client.post("/api/calculate/", json={"report_date": "2026-09-25", "recalculate": True})
    with Session() as db:
        own = db.query(CDU).filter_by(portfolio_code="310138-1").one()
        rows = {p.instrument_category: p.market_value_current for p in db.query(PortfolioPosition).filter_by(
            cdu_id=own.id, position_date=date(2026, 9, 25), instrument_code=None)}
    # opening cash = AC balance 28077 of 24.09; deposits of 24.09 and 25.09 are assets of the portfolio
    assert rows["DEPOSIT"] == pytest.approx(7_180_000_000 + 6_020_000_000, rel=1e-3)
    assert rows["CASH"] == pytest.approx(104_611_478.69 - 6_020_000_000)
    cash = _cash(client, own.id, "2026-09-25")["cash"]
    assert cash["opening"]["date"] == "2026-09-24" and cash["opening"]["total"] == pytest.approx(104_611_478.69)


# ───────────── п.9: TONIA Compounded ─────────────

@pytest.mark.skipif(not TONIA.exists(), reason="файл TONIA Compounded недоступен")
def test_tonia_import_is_idempotent_and_detected_in_documents(stand):
    client, Session, _ = stand
    r = client.post("/api/tonia/upload", files={"file": (TONIA.name, TONIA.read_bytes())})
    assert r.status_code == 200, r.text
    assert (r.json()["dates"], r.json()["inserted"]) == (1, 4)
    r = client.post("/api/primary-data/upload", files={"file": (TONIA.name, TONIA.read_bytes())})
    assert r.status_code == 200, r.text
    with Session() as db:
        assert db.query(ToniaCompounded).count() == 4
        v = {t.indicator: t.value for t in db.query(ToniaCompounded)}
    assert v == {"TCI": pytest.approx(2.2739959262), "TCR_1M": pytest.approx(16.1091),
                 "TCR_3M": pytest.approx(16.456), "TCR_6M": pytest.approx(17.499)}
    later = client.get("/api/tonia/value", params={"indicator": "TCR_3M", "date": "2026-10-20"}).json()
    assert later["date"] == "2026-10-14" and later["carried"] is True
    assert client.get("/api/tonia/value", params={"indicator": "TCR_3M", "date": "2026-10-13"}).status_code == 404
    rows = client.get("/api/tonia/").json()["rows"]
    assert rows[0]["TCR_6M"] == pytest.approx(17.499)


def test_floating_coupon_uses_tcr_plus_margin(stand):
    client, Session, _ = stand
    from services.automation.coupon_redemption_engine import run_daily_auto_events
    from services.market.tonia import upsert
    tansar = _cdu(Session, "Tansar Capital")
    with Session() as db:
        upsert(db, [{"date": date(2026, 1, 20), "TCR_3M": 16.0}, {"date": date(2026, 4, 20), "TCR_3M": 17.0}])
        _held_bond(db, tansar, isin="KZ2C0000FLT1", ticker="KFUSb99", face=1_000_000_000.0,
                   maturity=date(2028, 1, 21), coupon=15.0, freq=4, start=date(2025, 1, 21))
        db.flush()
        ref = db.query(InstrumentReference).filter_by(isin="KZ2C0000FLT1").one()
        ref.coupon_type, ref.float_index, ref.float_margin_pct = "float", "TCR_3M", 1.5
        db.commit()
        # without a schedule: regular quarterly periods; 21.04–21.07: TCR_3M on 21.04 =
        # the last value published before it (20.04) = 17.0 → 17.0 + 1.5 = 18.5 %
        res = run_daily_auto_events(db, date(2026, 7, 21))
        assert res["coupons"]["trades"] == 1
        ar = db.query(AccountReceivable).filter_by(isin="KZ2C0000FLT1").one()
        assert ar.amount == pytest.approx(1_000_000_000 * 0.185 * 90 / 360)
        # KASE schedule: periods without a rate get TCR + margin, a fixed rate is kept
        for d, rate in ((date(2026, 4, 21), None), (date(2026, 7, 21), None), (date(2026, 10, 21), 20.0)):
            db.add(CouponSchedule(isin="KZ2C0000FLT1", payment_date=d, rate_pct=rate, source="KASE"))
        db.commit()
        run_daily_auto_events(db, date(2026, 7, 21))
        db.refresh(ar)
        assert ar.amount == pytest.approx(1_000_000_000 * 0.185 * 90 / 360)
        run_daily_auto_events(db, date(2026, 10, 21))
        ar2 = db.query(AccountReceivable).filter_by(isin="KZ2C0000FLT1", record_date=date(2026, 10, 21)).one()
        assert ar2.amount == pytest.approx(1_000_000_000 * 0.20 * 90 / 360)


# ───────────── п.1: XLSX export with instrument sheets ─────────────

@pytest.mark.skipif(not RR_3112.exists(), reason="Risk Report 31.12.2025 недоступен")
def test_export_has_instrument_sheets_matching_summary(stand):
    client, Session, _ = stand
    from services.import_rr.risk_report_importer import import_risk_report
    with Session() as db:
        assert not import_risk_report(db, RR_3112, skip_if_imported=False).error
        db.commit()
    client.post("/api/calculate/", json={"report_date": "2025-12-31"})
    r = client.get("/api/export/xlsx", params={"report_date": "2025-12-31"})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content), data_only=True)
    for name in ("ГЦБ", "Агентские", "МФО", "РЕПО", "Дебиторская задолженность"):
        assert name in wb.sheetnames
    with Session() as db:
        summary = {}
        for p in db.query(PortfolioPosition).filter_by(position_date=date(2025, 12, 31), instrument_code=None):
            cdu = db.get(CDU, p.cdu_id)
            summary[(cdu.short_name or cdu.name, p.instrument_category)] = p.market_value_current
    for sheet, cat, col in (("ГЦБ", "GOV_BONDS", 15), ("МФО", "MFO_BONDS", 15), ("РЕПО", "REVERSE_REPO", 12)):
        totals = {str(row[0])[6:]: row[col - 1] for row in wb[sheet].iter_rows(min_row=5, values_only=True)
                  if row[0] and str(row[0]).startswith("Итого ")}
        assert totals, sheet
        for name, value in totals.items():
            assert value == pytest.approx(summary[(name, cat)], rel=1e-9), (sheet, name)
    halyk_gov = [row for row in wb["ГЦБ"].iter_rows(min_row=5, values_only=True) if row[0] and "Halyk" in str(row[0])
                 and not str(row[0]).startswith("Итого")]
    assert halyk_gov and all(r[2] for r in halyk_gov)          # ISIN in every row


# ───────────── review of 19.10.2026: edge cases of the events ─────────────

def _tansar_coupon_bond(db, tansar):
    _held_bond(db, tansar, isin="KZKD00001160", ticker="MUM096_0012", face=3_750_000_000.0,
               maturity=date(2029, 7, 28), coupon=10.55, freq=1, start=date(2021, 7, 28))


def _receivables(client, cdu_id: int, on: str) -> float:
    return client.get("/api/dashboard/category-breakdown",
                      params={"cdu_id": cdu_id, "category": "RECEIVABLES", "date": on}).json()["total"]


def test_moved_coupon_date_does_not_leave_a_second_coupon(stand):
    client, Session, _ = stand
    tansar = _cdu(Session, "Tansar Capital")
    with Session() as db:
        _tansar_coupon_bond(db, tansar)
        db.commit()
    client.post("/api/calculate/", json={"report_date": "2026-07-31", "recalculate": True})
    assert _receivables(client, tansar.id, "2026-07-31") == pytest.approx(395_625_000)
    with Session() as db:       # KASE schedule: the payment is shifted to 29.07
        for d in (date(2025, 7, 28), date(2026, 7, 29)):
            db.add(CouponSchedule(isin="KZKD00001160", payment_date=d, rate_pct=10.55, source="KASE"))
        db.commit()
    client.post("/api/calculate/", json={"report_date": "2026-07-31", "recalculate": True})
    with Session() as db:
        active = [(a.record_date, a.is_active) for a in db.query(AccountReceivable).filter_by(is_active=True)]
    assert [d for d, _ in active] == [date(2026, 7, 29)]
    assert _receivables(client, tansar.id, "2026-07-31") == pytest.approx(3_750_000_000 * 0.1055 * 361 / 360)


def test_receivable_confirmed_by_a_later_risk_report_is_kept(stand):
    client, Session, _ = stand
    tansar = _cdu(Session, "Tansar Capital")
    with Session() as db:
        _tansar_coupon_bond(db, tansar)
        db.commit()
    client.post("/api/calculate/", json={"report_date": "2026-07-28", "recalculate": True})
    from services.import_rr.risk_report_importer import _upsert_ar
    with Session() as db:
        doc = _rr_doc(db, date(2026, 7, 31))
        db.add(BondLot(cdu_id=tansar.id, isin="KZKD00001160", instrument_code="MUM096_0012", category="GOV_BONDS",
                       trade_date=date(2026, 3, 18), valuation_date=date(2026, 7, 31), face_value_current=3.75e9,
                       market_price=99.0, accrued_interest=0.0, total_value=3.75e9 * 0.99, source_doc_id=doc.id))
        cdu = db.get(CDU, tansar.id)
        _upsert_ar(db, [{"cdu_name": cdu.name, "isin": "KZKD00001160", "record_date": date(2026, 7, 28),
                         "balance_kzt": 395_625_000.0, "description": "Купон в пути"}],
                   {cdu.name: cdu.id}, doc.id, date(2026, 7, 31))
        db.commit()
    client.post("/api/calculate/", json={"report_date": "2026-08-03", "recalculate": True})
    with Session() as db:
        rec = db.query(PortfolioPosition).filter_by(cdu_id=tansar.id, position_date=date(2026, 8, 3),
                                                    instrument_code=None, instrument_category="RECEIVABLES").one()
        assert rec.market_value_current == pytest.approx(395_625_000)


def test_switch_to_cash_keeps_coupon_with_recorded_receipt(stand):
    client, Session, current = stand
    tansar = _cdu(Session, "Tansar Capital")
    with Session() as db:
        _tansar_coupon_bond(db, tansar)
        db.commit()
    client.post("/api/calculate/", json={"report_date": "2026-07-28", "recalculate": True})
    with Session() as db:
        ar_id = db.query(AccountReceivable).filter_by(cdu_id=tansar.id).one().id
    tx = client.post("/api/accounting/operations", json={
        "purpose": "RECEIVABLE_CLOSE", "cdu_id": tansar.id, "receivable_id": ar_id,
        "operation_date": "2026-07-30", "amount": "395625000"}).json()["created"][0]
    current["user"] = "checker"
    client.post(f"/api/accounting/{tx['entity']}/{tx['id']}/approve", json={})
    current["user"] = "maker"
    client.put("/api/formulas/parameters/coupon_settlement", json={"value": "cash", "reason": "тест"})
    try:
        client.post("/api/calculate/", json={"report_date": "2026-07-31", "recalculate": True})
        cash = _cash(client, tansar.id, "2026-07-31")["cash"]
        assert sum(m["amount_kzt"] for m in cash["movements"]) == pytest.approx(395_625_000)   # once
    finally:
        client.post("/api/formulas/parameters/coupon_settlement/reset", json={})


def test_float_settings_can_be_cleared_and_kase_amount_is_kept(stand):
    client, Session, _ = stand
    from services.portfolio.reference import ReferenceBook
    from services.market.tonia import upsert
    r = client.post("/api/securities/reference", json={"isin": "KZ2C0000FLT2", "ticker_kase": "KFUSb98",
                                                        "coupon_type": "float", "float_index": "TCR_3M",
                                                        "float_margin_pct": "1,5"})
    assert r.status_code == 200, r.text
    with Session() as db:
        ref = db.query(InstrumentReference).filter_by(isin="KZ2C0000FLT2").one()
        assert (ref.float_index, ref.float_margin_pct) == ("TCR_3M", 1.5)
        upsert(db, [{"date": date(2026, 1, 1), "TCR_3M": 16.0}])
        db.add(CouponSchedule(isin="KZ2C0000FLT2", payment_date=date(2026, 4, 21), period_start=date(2026, 1, 21),
                              amount_per_unit=42.0, source="KASE"))
        db.add(CouponSchedule(isin="KZ2C0000FLT2", payment_date=date(2026, 7, 21), source="KASE"))
        db.commit()
        rows = ReferenceBook(db).schedule("KZ2C0000FLT2")
        assert rows[0].rate is None and rows[0].amount_per_unit == 42.0       # KASE amount kept
        assert rows[1].rate == pytest.approx(0.175)                          # 16.0 + 1.5
    r = client.post("/api/securities/reference", json={"isin": "KZ2C0000FLT2", "float_index": "",
                                                        "float_margin_pct": None})
    assert r.status_code == 200, r.text
    with Session() as db:
        ref = db.query(InstrumentReference).filter_by(isin="KZ2C0000FLT2").one()
        assert (ref.float_index, ref.float_margin_pct) == (None, None)
