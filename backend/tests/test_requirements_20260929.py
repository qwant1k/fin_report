"""Regression tests for the testers' requirement log of 29.09.2026.

Real Trade Reports of the log (UD Capital 24.09, Alatau City Invest 16.09 and
24.09) are read from ``Примеры/Журнал 29.09.2026`` when available.
Managers are renamed to their legal names, as on the testers' stand.
"""
from __future__ import annotations

import io
import json
from datetime import date, timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import database as db_module
from auth import require_user, require_write
from database import Base, get_db
from models.db_models import (
    CashMovement, CDU, CDULimit, CouponSchedule, InstrumentReference, KaseRiskPrice,
    MBMIndex, PortfolioPosition, PortfolioSummary, RepoLot, Trade, User,
)

# «Примеры» lies inside the repository or next to it (C:\Projects\Сиуа\KDIF_FIN\Примеры)
LOG_DIR = next((p / "Примеры" for p in Path(__file__).resolve().parents[2:4] if (p / "Примеры").exists()),
               Path(__file__).resolve().parents[2] / "Примеры") / "Журнал 29.09.2026"
UD_24 = LOG_DIR / "Trade report UD от 24.09.2026.xlsx"
ACI_16 = LOG_DIR / "Trade report_ACI_16 09 2026.xlsx"
ACI_24 = LOG_DIR / "Trade report_ACI_24.09.2026.xlsx"
needs_files = pytest.mark.skipif(not ACI_16.exists(), reason="примеры журнала 29.09.2026 недоступны")

RENAME = {
    "Halyk Finance": 'АО "Дочерняя организация Народного Банка Казахстана "Halyk Finance"',
    "Centras Securities": 'АО "Сентрас Секьюритиз"',
    "Alatau City Invest": 'АО "Alatau City Invest"',
}


@pytest.fixture()
def stand(monkeypatch, tmp_path):
    engine = create_engine("sqlite:///:memory:", future=True,
                           connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Session = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
    monkeypatch.setattr(db_module, "engine", engine)
    monkeypatch.setattr(db_module, "SessionLocal", Session)
    Base.metadata.create_all(bind=engine)
    from seed import _seed_cdus, _seed_roles
    with Session() as db:
        _seed_roles(db)
        _seed_cdus(db)
        for cdu in db.query(CDU).all():
            if cdu.name in RENAME:
                cdu.name = RENAME[cdu.name]
        db.add(User(username="maker", password_hash="x", role="admin", is_active=True))
        db.add(User(username="checker", password_hash="x", role="admin", is_active=True))
        db.commit()
    from config import settings as app_settings
    monkeypatch.setattr(app_settings, "upload_dir", str(tmp_path))
    from main import create_app
    app = create_app()
    current = {"user": "maker"}

    def _db():
        s = Session()
        try:
            yield s
        finally:
            s.close()

    def _user():
        with Session() as s:
            return s.query(User).filter_by(username=current["user"]).first()

    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[require_user] = _user
    app.dependency_overrides[require_write] = _user
    client = TestClient(app)
    yield client, Session, current
    engine.dispose()


def _cdu(Session, canonical: str) -> CDU:
    from services.cdu_lookup import find_cdu
    with Session() as db:
        return find_cdu(db, name=canonical)


def _upload(client, path: Path):
    r = client.post("/api/upload/trade-report", files={"file": (path.name, path.read_bytes())})
    assert r.status_code == 200, r.text
    return r.json()


# ───────────── 1, 5, 6, 9: import of UD / Alatau, cash on settlement dates ─────────────

@needs_files
def test_trade_reports_of_renamed_managers_are_imported(stand):
    client, Session, _ = stand
    for path in (UD_24, ACI_16, ACI_24):
        _upload(client, path)
    ud = _cdu(Session, "UD Capital")
    aci = _cdu(Session, "Alatau City Invest")
    with Session() as db:
        n_ud = db.query(Trade).filter_by(cdu_id=ud.id, is_active=True).count()
        n_aci = db.query(Trade).filter_by(cdu_id=aci.id, is_active=True).count()
        assert n_ud == 8          # 4 REPO × (open + close)
        assert n_aci == 18 + 1 + 4  # 9 REPO on 16.09, purchase NTK, 2 REPO on 24.09
        by_date = {}
        for cm in db.query(CashMovement).filter_by(cdu_id=aci.id, is_active=True):
            by_date[cm.operation_date] = by_date.get(cm.operation_date, 0.0) + cm.amount_kzt
    # purchase of 16.09 settles on 18.09 (T+2)
    assert by_date[date(2026, 9, 18)] == pytest.approx(-999_999_980.72)
    assert by_date[date(2026, 9, 16)] < 0 < by_date[date(2026, 9, 17)]


@needs_files
def test_purchase_is_in_portfolio_from_settlement_date_and_positions_show_it(stand):
    client, Session, _ = stand
    _upload(client, ACI_16)
    aci = _cdu(Session, "Alatau City Invest")
    p17 = client.get("/api/positions/as-of", params={"as_of": "2026-09-17", "cdu_id": aci.id}).json()
    assert not [r for r in p17["rows"] if r["instrument_category"] == "GOV_BONDS"]
    p18 = client.get("/api/positions/as-of", params={"as_of": "2026-09-18", "cdu_id": aci.id}).json()
    note = [r for r in p18["rows"] if r["instrument_category"] == "GOV_BONDS"]
    assert len(note) == 1
    assert note[0]["net_quantity"] == pytest.approx(10_396_988)
    assert note[0]["market_value"] == pytest.approx(999_999_980.72, rel=1e-6)
    # a discount note without a quote accretes at the purchase yield
    p25 = client.get("/api/positions/as-of", params={"as_of": "2026-09-25", "cdu_id": aci.id}).json()
    later = [r for r in p25["rows"] if r["instrument_category"] == "GOV_BONDS"][0]
    assert later["market_value"] > note[0]["market_value"]
    assert later["duration"] and later["duration"] > 0


@needs_files
def test_repo_legs_are_linked_by_head_deal_even_when_close_comes_first(stand):
    client, Session, _ = stand
    _upload(client, ACI_16)
    aci = _cdu(Session, "Alatau City Invest")
    with Session() as db:
        lots = db.query(RepoLot).filter_by(cdu_id=aci.id).all()
    assert len(lots) == 9
    assert all(l.open_trade_id and l.close_trade_id for l in lots)


@needs_files
def test_category_drilldown_lists_repo_of_the_date(stand):
    client, Session, _ = stand
    _upload(client, UD_24)
    ud = _cdu(Session, "UD Capital")
    client.post("/api/calculate/", json={"report_date": "2026-09-24", "recalculate": True})
    d = client.get("/api/dashboard/instrument-details",
                   params={"cdu_id": ud.id, "category": "REVERSE_REPO", "from": "2026-09-24", "to": "2026-09-24"}).json()
    assert len(d["rows"]) == 4
    assert sum(r["amount"] for r in d["rows"]) == pytest.approx(1_297_349_973.98, abs=0.05)


# ───────────── 4: reverse REPO accrues daily ─────────────

def test_open_reverse_repo_value_grows_by_daily_reward(stand):
    _, Session, _ = stand
    from services.portfolio.engine import build_portfolio
    with Session() as db:
        cdu = db.query(CDU).filter_by(short_name="Halyk").first()
        t = Trade(cdu_id=cdu.id, trade_date=date(2026, 9, 24), value_date=date(2026, 9, 24),
                  operation_type="REPO_OPEN", deal_id="1", amount_kzt=-1_000_000_000.0,
                  is_active=True, is_manual=False, workflow_status="APPROVED")
        db.add(t)
        db.flush()
        db.add(RepoLot(cdu_id=cdu.id, instrument_code="MUM180_0011", deal_id="1",
                       trade_date=date(2026, 9, 24), valuation_date=date(2026, 9, 24),
                       close_date=date(2026, 10, 1), term_days=7,
                       face_value=1_000_000_000.0, close_value=1_003_000_000.0,
                       repo_rate_pct=15.6, open_trade_id=t.id))
        db.commit()
        v24 = build_portfolio(db, cdu.id, date(2026, 9, 24)).repos[0].market_value
        v25 = build_portfolio(db, cdu.id, date(2026, 9, 25)).repos[0].market_value
    assert v24 == pytest.approx(1_000_000_000.0)
    assert v25 - v24 == pytest.approx(3_000_000.0 / 7, rel=1e-6)


# ───────────── 3: benchmark = MBM ModDuration of the date ─────────────

def test_benchmark_is_mbm_mod_duration_of_the_report_date(stand):
    client, Session, _ = stand
    from services.calculator.portfolio_calculator import _benchmark_for
    with Session() as db:
        db.add_all([
            MBMIndex(index_date=date(2026, 9, 24), ytm_value=1316.63, duration=2.64, mod_duration=2.31),
            MBMIndex(index_date=date(2026, 9, 28), ytm_value=1318.90, duration=2.63, mod_duration=2.30),
        ])
        db.commit()
        assert _benchmark_for(db, date(2026, 9, 24)) == pytest.approx(2.31)
        assert _benchmark_for(db, date(2026, 9, 26)) == pytest.approx(2.31)
        assert _benchmark_for(db, date(2026, 9, 28)) == pytest.approx(2.30)
    payload = client.get("/api/dashboard/summary", params={"report_date": "2026-09-24"}).json()
    assert payload["benchmark_duration"] == pytest.approx(2.31)


# ───────────── 7: limits of the reference book in force on the date ─────────────

def test_limits_of_the_reference_book_in_force_on_the_date(stand):
    _, Session, _ = stand
    from services.calculator.limit_checker import check_limit_spec
    from services.calculator.portfolio_calculator import _load_limit_specs
    with Session() as db:
        cdu = db.query(CDU).filter_by(short_name="Halyk").first()
        db.query(CDULimit).filter_by(cdu_id=cdu.id).delete()
        db.add_all([
            CDULimit(cdu_id=cdu.id, instrument_category="REVERSE_REPO", min_limit_pct=0.0,
                     max_limit_pct=0.5, hard_limit_pct=0.5, soft_limit_pct=0.45,
                     valid_from=date(2026, 1, 1), valid_to=date(2026, 9, 24)),
            CDULimit(cdu_id=cdu.id, instrument_category="REVERSE_REPO", min_limit_pct=0.0,
                     max_limit_pct=0.7, hard_limit_pct=0.7, soft_limit_pct=0.6,
                     valid_from=date(2026, 9, 25)),
        ])
        db.commit()
        old = _load_limit_specs(db, cdu.id, date(2026, 9, 24))["REVERSE_REPO"]
        new = _load_limit_specs(db, cdu.id, date(2026, 9, 25))["REVERSE_REPO"]
    assert (old.max_pct, new.max_pct) == (0.5, 0.7)
    assert check_limit_spec(0.47, old) == ("ok", "breach")     # soft warning only
    assert check_limit_spec(0.55, old) == ("breach", "breach")
    assert check_limit_spec(0.55, new) == ("ok", "ok")


# ───────────── 8: instrument of the reference book without a manager ─────────────

def test_security_can_be_saved_without_a_manager(stand):
    client, Session, _ = stand
    r = client.post("/api/securities/", json={
        "isin": "KZKD00001087", "instrument_code": "MUM108_0013", "category": "GOV_BONDS",
        "instrument_name": "МЕУКАМ-108", "nominal_per_unit": 1000, "coupon_rate_pct": 11,
        "maturity_date": "2029-03-31", "start_date": "2020-03-31", "frequency": 1, "base": "30/360",
    })
    assert r.status_code == 200, r.text
    assert r.json()["cdu_id"] is None
    with Session() as db:
        ref = db.query(InstrumentReference).filter_by(isin="KZKD00001087").one()
        assert (ref.ticker_kase, ref.nominal, ref.frequency) == ("MUM108_0013", 1000, 1)
    # the same instrument without a manager cannot be duplicated
    assert client.post("/api/securities/", json={"isin": "KZKD00001087"}).status_code == 409


# ───────────── 10: coupon schedule — import, manual edit, used in calculations ─────────────

def _kase_schedule_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Дата начала купонной выплаты", "Ставка, % год", "Дата фиксации реестра", "Дата окончания выплаты"])
    for i, (d, end) in enumerate([("31.03.21", "31.03.21"), ("31.03.22", "31.03.22"), ("31.03.23", "31.03.23"),
                                  ("31.03.24", "01.04.24"), ("31.03.25", "31.03.25"), ("31.03.26", "31.03.26"),
                                  ("31.03.27", "31.03.27"), ("31.03.28", "31.03.28"), ("31.03.29", "31.03.29")], 1):
        ws.append([i, d, "11,00", "–", end])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_coupon_schedule_import_edit_and_use(stand):
    client, Session, _ = stand
    client.post("/api/securities/reference", json={
        "isin": "KZKD00001087", "ticker_kase": "MUM108_0013", "nominal": 1000, "coupon_rate_pct": 11,
        "frequency": 1, "base": "30/360", "start_date": "2020-03-31", "maturity_date": "2029-03-31",
    })
    r = client.post("/api/securities/reference/coupons/import",
                    files={"file": ("graf.xlsx", _kase_schedule_xlsx())}, data={"isin": "MUM108_0013"})
    assert r.status_code == 200, r.text
    assert r.json()["imported"] == {"KZKD00001087": 9}
    sched = client.get("/api/securities/reference/KZKD00001087/coupons").json()["rows"]
    assert sched[3]["payment_end_date"] == "2024-04-01"
    # manual correction: floating rate from 2027
    rows = sched
    for row in rows:
        if row["payment_date"] >= "2027-01-01":
            row["rate_pct"] = 12.5
    assert client.put("/api/securities/reference/KZKD00001087/coupons", json={"rows": rows}).status_code == 200

    from services.portfolio import bond_math as bm
    from services.portfolio import coupon_schedule as cs
    from services.portfolio.reference import ReferenceBook
    with Session() as db:
        refs = ReferenceBook(db)
        schedule = refs.schedule("MUM108_0013")
        assert len(schedule) == 9
        # НКД on 24.09.2026: the period 31.03.26 – 31.03.27 is paid on 31.03.27 at the
        # edited rate 12.5 %; 30E/360 days 31.03.26 → 24.09.26 = 174
        acc = cs.accrued(1_000_000.0, schedule, date(2026, 9, 24), bm.BASIS_30E_360)
        assert acc == pytest.approx(1_000_000 * 0.125 * 174 / 360)
        acc_2025 = cs.accrued(1_000_000.0, schedule, date(2025, 9, 24), bm.BASIS_30E_360)
        assert acc_2025 == pytest.approx(1_000_000 * 0.11 * 174 / 360)
        # the 2027 coupon uses the edited rate
        amount, start, rate = cs.coupon_on(schedule, date(2027, 3, 31), 1_000_000.0, bm.BASIS_30E_360)
        assert (start, rate) == (date(2026, 3, 31), pytest.approx(0.125))
        assert amount == pytest.approx(125_000.0)


def test_coupon_schedule_can_be_generated_from_issue_parameters(stand):
    client, _, _ = stand
    client.post("/api/securities/reference", json={
        "isin": "KZKD00001087", "nominal": 1000, "coupon_rate_pct": 11, "frequency": 1,
        "start_date": "2020-03-31", "maturity_date": "2029-03-31",
    })
    rows = client.post("/api/securities/reference/KZKD00001087/coupons/generate").json()["rows"]
    assert [r["payment_date"] for r in rows][:2] == ["2021-03-31", "2022-03-31"]
    assert rows[-1]["payment_date"] == "2029-03-31" and len(rows) == 9


# ───────────── 11: KASE prices calculated from risk parameters ─────────────

def _risk_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Расчетные цены_ФР"
    ws.append(["Дата оценки"])
    ws.append([date(2026, 9, 24)])
    ws.append([])
    ws.append(["№ п/п", "Торговый код", "ISIN", "Алгоритмы и параметры расчета ЦБ", "Краткое наименование эмитента",
               "Расчетная цена", "Расчетная грязная", "Доходность до погашения. % годовых", "Дней до погашения",
               "Единица измерения цены"])
    ws.append([1, "NTK092_2889", "KZW100000000", "A1", "НБРК", 96.43, 96.43, 16.1, 84, "%"])
    wb.create_sheet("Расчетные цены_фьюч.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_risk_parameter_prices_are_loaded_from_xlsx_file_with_date_and_source(stand):
    client, Session, _ = stand
    r = client.post("/api/kase/risk/upload", files={"file": ("risk_24092026.xlsx", _risk_xlsx())},
                    data={"report_date": "2026-09-24", "propagate": "false"})
    assert r.status_code == 200, r.text
    with Session() as db:
        row = db.execute(select(KaseRiskPrice)).scalars().one()
    assert (row.trade_date, row.valuation_date) == (date(2026, 9, 24), date(2026, 9, 24))
    assert row.settlement_price == pytest.approx(96.43)
    assert row.source == "file:risk_24092026.xlsx"


# ───────────── 12: manual cash operations with account and purpose ─────────────

def test_manual_cash_receipt_with_account_changes_cash_after_approval(stand):
    client, Session, current = stand
    aci = _cdu(Session, "Alatau City Invest")
    r = client.post("/api/accounting/cash-movements", json={
        "cdu_id": aci.id, "account": "kz72 8562 2021 3919 6235", "operation_date": "2026-09-15",
        "currency": "KZT", "amount": "1500000000", "movement_type": "RECEIPT",
        "basis": "Перечисление средств КФГД в доверительное управление",
        "change_reason": "Поступление отсутствует в Trade Report",
    })
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["account"] == "KZ728562202139196235"
    assert body["amount"] == 1_500_000_000
    # four-eyes: the author cannot approve, another user can
    current["user"] = "checker"
    ok = client.post(f"/api/accounting/cash-movements/{body['id']}/approve", json={})
    assert ok.status_code == 200, ok.text
    client.post("/api/calculate/", json={"report_date": "2026-09-15", "recalculate": True})
    with Session() as db:
        cash = db.execute(select(PortfolioPosition).where(
            PortfolioPosition.cdu_id == aci.id, PortfolioPosition.position_date == date(2026, 9, 15),
            PortfolioPosition.instrument_category == "CASH",
        )).scalars().one()
    assert cash.market_value_current == pytest.approx(1_500_000_000)
    accounts = client.get("/api/accounting/accounts", params={"cdu_id": aci.id}).json()
    assert accounts[0]["account"] == "KZ728562202139196235"


def test_payment_types_get_a_negative_sign(stand):
    client, Session, _ = stand
    aci = _cdu(Session, "Alatau City Invest")
    r = client.post("/api/accounting/cash-movements", json={
        "cdu_id": aci.id, "operation_date": "2026-09-15", "currency": "KZT", "amount": "1000",
        "movement_type": "COMMISSION", "basis": "Комиссия кастодиана", "change_reason": "Ручной ввод",
    })
    assert r.json()["amount"] == -1000


# ───────────── CDU lookup ─────────────

def test_manager_lookup_by_kase_codes_and_legal_names(stand):
    _, Session, _ = stand
    from services.cdu_lookup import find_cdu
    with Session() as db:
        assert find_cdu(db, prefix="DRBES").short_name == "Alatau"
        assert find_cdu(db, prefix="ACI").short_name == "Alatau"
        assert find_cdu(db, prefix="CAIFC").short_name == "Tansar"
        assert find_cdu(db, name="Сентрас").short_name == "Centras"
        assert find_cdu(db, name="Halyk Finance").short_name == "Halyk"
        assert find_cdu(db, name="Study Invest") is None
