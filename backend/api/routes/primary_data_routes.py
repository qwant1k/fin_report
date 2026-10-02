"""API endpoints for primary-data import from custodians (ЧДУ / НБ РК)."""
from __future__ import annotations

import hashlib
import json
import re
import shutil
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, UploadFile
from loguru import logger
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from auth import require_permission, require_user
from config import settings
from database import SessionLocal, get_db
from models.db_models import CDU, ImportJob, SourceDocument, User
from services.calculator.constants import normalize_cdu_name
from services.import_primary.binding import bind_certificate
from services.import_primary.cert_parser import parse_certificate
from services.import_primary.holdings_parser import import_holdings_xlsx
from services.import_primary.recon_parser import parse_reconciliation_xlsx
from services.import_primary.statement_parser import parse_pdf_statement
from services.audit import write_audit
from services.parser.trade_importer import import_single_trade_report_xlsx
from services.parser.trade_ingest import ingest_trade_report

router = APIRouter(
    prefix="/api/primary-data",
    tags=["primary-data"],
    dependencies=[Depends(require_permission("page.primary_data"))],
)


class ImportResponse(BaseModel):
    job_id: int
    status: str
    message: str
    counters: dict = Field(default_factory=dict)


@router.post("/upload", response_model=ImportResponse)
async def upload_primary_data(
    file: UploadFile = File(...),
    cdu_id: Optional[int] = Form(None),
    doc_date: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("primary_data.upload")),
):
    """Upload one primary-data file, save it for audit, and import immediately.

    ``cdu_id`` / ``doc_date`` are optional overrides; otherwise the manager and
    the date are recognised from the content (statements: by the account
    number seen in earlier statements of the manager).
    """
    original_name = file.filename or "unnamed"
    safe_name = original_name.replace("/", "_").replace("\\", "_")
    staging = settings.upload_path / "primary_data" / "_incoming"
    staging.mkdir(parents=True, exist_ok=True)
    target = staging / f"{datetime.utcnow().strftime('%Y%m%d_%H%M%S_%f')}_{safe_name}"
    with target.open("wb") as out:
        shutil.copyfileobj(file.file, out)

    file_type = _detect_file_type(original_name, target)
    if file_type == "unknown":
        target.unlink(missing_ok=True)
        raise HTTPException(400, f"Не удалось определить тип документа «{original_name}» "
                                 "(ожидаются Trade Report, Holdings, сверка ЧДУ, биржевое свидетельство или выписка)")
    final_dir = settings.upload_path / "primary_data" / file_type
    final_dir.mkdir(parents=True, exist_ok=True)
    final_path = final_dir / target.name
    shutil.move(str(target), final_path)
    target = final_path

    override_date: Optional[date] = None
    if doc_date:
        try:
            override_date = date.fromisoformat(doc_date)
        except ValueError:
            raise HTTPException(400, "doc_date должен быть в формате YYYY-MM-DD")

    content = target.read_bytes()
    sha = hashlib.sha256(content).hexdigest()
    duplicate = db.execute(
        select(SourceDocument).where(
            SourceDocument.sha256 == sha,
            SourceDocument.parse_status == "OK",
        )
    ).scalars().first()
    if duplicate:
        return ImportResponse(
            job_id=-1,
            status="SKIPPED",
            message=f"Документ уже загружен ранее: id={duplicate.id} ({duplicate.file_name})",
        )

    job = ImportJob(
        job_type="FOLDER_PRIMARY",
        status="RUNNING",
        triggered_by=user.username,
        files_total=1,
        params_json=json.dumps({"file_name": original_name, "file_type": file_type}, ensure_ascii=False),
    )
    db.add(job)
    db.flush()

    try:
        if file_type == "trade_report":
            ingest = ingest_trade_report(
                db, path=target, original_name=original_name, actor=user.username,
                cdu_id=cdu_id, trade_date=override_date,
            )
            counters = {
                "cdu_name": ingest.cdu.name if ingest.cdu else None,
                "trade_date": ingest.trade_date.isoformat(),
                "trade_file_id": ingest.trade_file.id,
                **ingest.counters,
                "warnings": ingest.parsed.warnings,
            }
            source_doc = ingest.source_doc
        else:
            source_doc = SourceDocument(
                doc_type=_to_doc_type(file_type),
                cdu_id=None,
                doc_date=None,
                file_name=original_name,
                file_path=str(target),
                sha256=sha,
                file_size=len(content),
                uploaded_by=user.username,
                parse_status="PENDING",
            )
            db.add(source_doc)
            db.flush()
            counters = _run_import(db, str(target), file_type, source_doc.id, user.username)
            _finalize_source_doc(db, source_doc, counters, cdu_id=cdu_id, doc_date=override_date)
        job.params_json = json.dumps(
            {"source_doc_id": source_doc.id, "file_name": original_name, "file_type": file_type},
            ensure_ascii=False,
        )
        job.status = "DONE"
        job.finished_at = datetime.utcnow()
        job.files_done = 1
        job.rows_imported = source_doc.rows_imported
        message = f"Импортирован документ «{_TYPE_LABELS.get(file_type, file_type)}»"
        if file_type == "nbrk_daily":
            message += (f": депозитов {counters.get('deposits', 0)}, РЕПО {counters.get('repos', 0)}, "
                        f"сделок с ЦБ {counters.get('securities', 0)}; остатки по {len(counters.get('accounts') or [])} счетам "
                        "сохранены для сверки")
            if counters.get("warnings"):
                message += f"; предупреждений: {len(counters['warnings'])}"
        elif source_doc.cdu_id is None:
            message += " — ЧДУ не определён, укажите его вручную"
        write_audit(db, user=user.username, action="PRIMARY_DOCUMENT_UPLOADED",
                    entity="SourceDocument", entity_id=source_doc.id,
                    details={"file_type": file_type, "file_name": original_name,
                             "cdu_id": source_doc.cdu_id,
                             "doc_date": source_doc.doc_date.isoformat() if source_doc.doc_date else None})
    except Exception as exc:
        logger.exception("Primary data import failed")
        db.rollback()
        job = ImportJob(
            job_type="FOLDER_PRIMARY", status="FAILED", triggered_by=user.username,
            files_total=1, files_failed=1, finished_at=datetime.utcnow(),
            log=str(exc)[:4000],
            params_json=json.dumps({"file_name": original_name, "file_type": file_type}, ensure_ascii=False),
        )
        db.add(job)
        counters = {}
        message = f"Ошибка импорта: {exc}"
    finally:
        db.commit()

    return ImportResponse(job_id=job.id, status=job.status, message=message,
                          counters=json.loads(json.dumps(counters, default=str)))


@router.put("/documents/{doc_id}")
def bind_document(
    doc_id: int,
    payload: dict,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("primary_data.upload")),
):
    """Manually bind a primary document to a CDU and/or date."""
    doc = db.get(SourceDocument, doc_id)
    if doc is None:
        raise HTTPException(404, "Документ не найден")
    before = {"cdu_id": doc.cdu_id, "doc_date": doc.doc_date.isoformat() if doc.doc_date else None}
    if payload.get("cdu_id") is not None:
        cdu = db.get(CDU, int(payload["cdu_id"]))
        if cdu is None:
            raise HTTPException(400, "ЧДУ не найден")
        doc.cdu_id = cdu.id
    if payload.get("doc_date"):
        try:
            doc.doc_date = date.fromisoformat(str(payload["doc_date"]))
        except ValueError:
            raise HTTPException(400, "doc_date должен быть в формате YYYY-MM-DD")
    write_audit(db, user=user.username, action="PRIMARY_DOCUMENT_BOUND",
                entity="SourceDocument", entity_id=doc.id,
                details={"before": before, "after": {"cdu_id": doc.cdu_id,
                         "doc_date": doc.doc_date.isoformat() if doc.doc_date else None}})
    db.commit()
    return {"ok": True, "cdu_id": doc.cdu_id, "doc_date": doc.doc_date.isoformat() if doc.doc_date else None}


@router.post("/bulk-folder")
async def bulk_import_primary_folder(
    payload: dict,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("primary_data.upload")),
):
    """Queue recursive folder import for primary-data packages."""
    folder_path = payload.get("folder_path")
    if not folder_path:
        raise HTTPException(400, "folder_path обязателен")
    folder = Path(folder_path)
    if not folder.exists() or not folder.is_dir():
        raise HTTPException(400, (
            f"Папка не найдена на сервере: {folder}. Путь указывается так, как его видит сервер "
            "приложения (при запуске в Docker — смонтированный каталог, например /data/...), "
            "а не путь на компьютере пользователя; иначе загрузите файлы через форму."
        ))

    job = ImportJob(
        job_type="FOLDER_PRIMARY",
        status="RUNNING",
        triggered_by=user.username,
        params_json=json.dumps({"folder_path": str(folder)}, ensure_ascii=False),
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    background_tasks.add_task(_bulk_primary_job, job.id, str(folder), user.username)
    return {"job_id": job.id, "status": job.status}


@router.get("/documents")
def list_source_documents(limit: int = 50, db: Session = Depends(get_db)):
    rows = db.execute(
        select(SourceDocument, CDU.name.label("cdu_name"))
        .outerjoin(CDU, SourceDocument.cdu_id == CDU.id)
        .order_by(SourceDocument.uploaded_at.desc())
        .limit(limit)
    ).all()
    return [
        {
            "id": row.SourceDocument.id,
            "filename": row.SourceDocument.file_name,
            "file_type": row.SourceDocument.doc_type,
            "file_size": row.SourceDocument.file_size,
            "import_status": row.SourceDocument.parse_status,
            "imported_at": row.SourceDocument.uploaded_at.isoformat() if row.SourceDocument.uploaded_at else None,
            "imported_by": row.SourceDocument.uploaded_by,
            "source_cdu": row.cdu_name,
            "doc_date": row.SourceDocument.doc_date.isoformat() if row.SourceDocument.doc_date else None,
            "rows_imported": row.SourceDocument.rows_imported,
        }
        for row in rows
    ]


_TYPE_LABELS = {
    "trade_report": "Trade Report",
    "holdings": "Holdings report",
    "reconciliation": "Сверка ЧДУ",
    "exchange_certificate": "Биржевое свидетельство",
    "pdf_statement": "Выписка по счёту",
    "nbrk_daily": "Файл НБ РК (КФГД_ГГГГММДД)",
}


def _detect_file_type(filename: str, path: Optional[Path] = None) -> str:
    """Document type by name, then by content (names are not reliable:
    «КФГД-100925_KZT.pdf» and «11092025 КФГД.pdf» are statements)."""
    f = filename.lower()
    from services.import_primary.nbrk_parser import looks_like_nbrk_file
    if f.endswith((".csv", ".txt")) and looks_like_nbrk_file(
        filename, path.read_bytes()[:4000] if path is not None and path.exists() else None,
    ):
        return "nbrk_daily"
    if any(k in f for k in ("trade report", "trade_report", "сделки", "trade")) and f.endswith((".xlsx", ".xlsm")):
        return "trade_report"
    if any(k in f for k in ("holdings", "позиции", "отчет об активах", "активах")) and f.endswith(".xlsx"):
        return "holdings"
    if any(k in f for k in ("сверка", "recon", "reconciliation")) and f.endswith(".xlsx"):
        return "reconciliation"
    if any(k in f for k in ("биржев", "свидетельств", "certificate", "cert")) and f.endswith((".pdf", ".png", ".jpg", ".jpeg", ".docx")):
        return "exchange_certificate"
    if any(k in f for k in ("выписка", "statement")) and f.endswith(".pdf"):
        return "pdf_statement"
    if path is None or not path.exists():
        return "unknown"
    try:
        if f.endswith(".pdf"):
            from services.import_primary.statement_parser import _extract_pdf_text, looks_like_statement
            text = _extract_pdf_text(path) or ""
            if looks_like_statement(text):
                return "pdf_statement"
            return "exchange_certificate"   # scanned PDFs without text are certificates in practice
        if f.endswith((".png", ".jpg", ".jpeg", ".docx")):
            return "exchange_certificate"
        if f.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb[wb.sheetnames[0]]
                head = " ".join(
                    str(v).lower() for row in ws.iter_rows(min_row=1, max_row=12, values_only=True)
                    for v in row if v is not None
                )
            finally:
                wb.close()
            if "к/п" in head and ("код инструмента" in head or "код режима" in head):
                return "trade_report"
            if "суммарно-текущая рыночная стоимость" in head or "итого активов" in head or "остатки на инвест" in head:
                return "reconciliation"
            if "isin" in head and ("количество" in head or "номинал" in head):
                return "holdings"
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"Content sniffing failed for {filename}: {exc!r}")
    return "unknown"


def _run_import(db: Session, file_path: str, file_type: str, source_doc_id: int | None, user: str) -> dict:
    if file_type == "trade_report":
        return import_single_trade_report_xlsx(db, file_path, uploaded_by=user, source_doc_id=source_doc_id)
    if file_type == "holdings":
        return import_holdings_xlsx(db, file_path, uploaded_by=user, source_doc_id=source_doc_id)
    if file_type == "reconciliation":
        return parse_reconciliation_xlsx(file_path)
    if file_type == "exchange_certificate":
        return parse_certificate(file_path)
    if file_type == "pdf_statement":
        return parse_pdf_statement(file_path)
    if file_type == "nbrk_daily":
        from services.import_primary.nbrk_parser import import_nbrk_file
        return import_nbrk_file(db, file_path, source_doc_id=source_doc_id, actor=user)
    raise ValueError(f"Unsupported file_type {file_type}")


def _to_doc_type(file_type: str) -> str:
    return {
        "trade_report": "TRADE_REPORT",
        "holdings": "HOLDINGS",
        "reconciliation": "RECONCILIATION",
        "exchange_certificate": "CERTIFICATE",
        "pdf_statement": "STATEMENT_PDF",
        "nbrk_daily": "NBRK_DAILY",
    }.get(file_type, "OTHER")


def _cdu_from_accounts(db: Session, accounts: list[str]) -> Optional[int]:
    """Statements carry the fund as client; the manager is identified by the
    account number already bound to a manager in an earlier statement."""
    if not accounts:
        return None
    docs = db.execute(select(SourceDocument).where(
        SourceDocument.doc_type == "STATEMENT_PDF",
        SourceDocument.cdu_id.is_not(None),
    ).order_by(SourceDocument.id.desc())).scalars().all()
    for doc in docs:
        try:
            meta = json.loads(doc.parse_meta_json or "{}")
        except ValueError:
            continue
        if set(meta.get("accounts_found") or []) & set(accounts):
            return doc.cdu_id
    return None


def _finalize_source_doc(
    db: Session, source_doc: SourceDocument, counters: dict,
    *, cdu_id: Optional[int] = None, doc_date: Optional[date] = None,
) -> None:
    source_doc.parse_status = "OK"
    source_doc.parsed_at = datetime.utcnow()
    source_doc.parse_meta_json = json.dumps(counters, ensure_ascii=False, default=str)
    source_doc.rows_imported = _count_rows(counters)

    if cdu_id:
        source_doc.cdu_id = cdu_id
    else:
        cdu_name = counters.get("cdu_name")
        if cdu_name:
            from services.cdu_lookup import find_cdu
            cdu = find_cdu(db, name=cdu_name)
            if cdu:
                source_doc.cdu_id = cdu.id
        if source_doc.cdu_id is None and source_doc.doc_type == "STATEMENT_PDF":
            source_doc.cdu_id = _cdu_from_accounts(db, counters.get("accounts_found") or [])

    report_date = doc_date or counters.get("trade_date") or counters.get("report_date")
    if hasattr(report_date, "isoformat"):
        source_doc.doc_date = report_date
    elif isinstance(report_date, str):
        try:
            source_doc.doc_date = datetime.fromisoformat(report_date).date()
        except ValueError:
            pass

    if source_doc.cdu_id is None and source_doc.doc_type == "CERTIFICATE":
        # KASE certificates name the fund, not the manager: bind by deal numbers
        bind_certificate(db, source_doc)


def _count_rows(counters: dict) -> int:
    total = 0
    for key in ("trades", "cash_snapshots", "portfolio_positions", "rows_parsed"):
        value = counters.get(key)
        if isinstance(value, int):
            total += value
    # certificates: recognised deals; statements: account sections
    for key in ("deals", "sections"):
        value = counters.get(key)
        if isinstance(value, list):
            total += len(value)
    return total


def _bulk_primary_job(job_id: int, folder_path: str, user: str) -> None:
    db = SessionLocal()
    try:
        job = db.get(ImportJob, job_id)
        if not job:
            return
        files = [p for p in sorted(Path(folder_path).rglob("*")) if p.is_file()]
        job.files_total = len(files)
        db.commit()

        total_rows = 0
        for path in files:
            file_type = _detect_file_type(path.name, path)
            if file_type == "unknown":
                continue
            try:
                if file_type == "trade_report":
                    ingest = ingest_trade_report(db, path=path, original_name=path.name, actor=user)
                    counters = dict(ingest.counters)
                else:
                    sha = hashlib.sha256(path.read_bytes()).hexdigest()
                    if db.execute(select(SourceDocument.id).where(
                        SourceDocument.sha256 == sha, SourceDocument.parse_status == "OK",
                    )).first():
                        job.files_done += 1
                        continue
                    doc = SourceDocument(
                        doc_type=_to_doc_type(file_type), file_name=path.name,
                        file_path=str(path), sha256=sha, file_size=path.stat().st_size,
                        uploaded_by=user, parse_status="PENDING",
                    )
                    db.add(doc)
                    db.flush()
                    counters = _run_import(db, str(path), file_type, doc.id, user)
                    _finalize_source_doc(db, doc, counters)
                total_rows += _count_rows(counters)
                job.files_done += 1
            except Exception as exc:
                job.files_failed += 1
                job.log = (job.log or "") + f"\n[ERR] {path}: {exc}"
                logger.warning(f"Failed to import {path.name}: {exc}")
            db.commit()

        job.rows_imported = total_rows
        job.status = "DONE" if job.files_failed == 0 else "PARTIAL"
        job.finished_at = datetime.utcnow()
        db.commit()
    except Exception as exc:
        logger.exception("Bulk primary import failed")
        job = db.get(ImportJob, job_id)
        if job:
            job.status = "FAILED"
            job.finished_at = datetime.utcnow()
            job.log = (job.log or "") + f"\n[FATAL] {exc}"
            db.commit()
    finally:
        db.close()
